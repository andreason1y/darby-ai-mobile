#!/usr/bin/env python3
"""Build a formal five-part legal opinion DOCX.

The opinion follows the standard law-firm structure:
  I.   Statement of Facts
  II.  Legal Issues Presented
  III. Regulatory & Precedent Framework
  IV.  Substantive Legal Analysis
  V.   Conclusions & Actionable Advice

All facts, authorities and conclusions come from the JSON the caller supplies.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Mapping, Sequence

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, ValidationError


NAVY = RGBColor(0x1F, 0x38, 0x64)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)
NAVY_FILL = "1F3864"
PALE_FILL = "EDF2FA"

MONTHS = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]


class OpinionError(Exception):
    """An opinion-data or document problem that should be shown to the caller."""


class Flexible(BaseModel):
    model_config = ConfigDict(extra="allow", populate_by_name=True)


class FirmInfo(Flexible):
    name: str = ""
    address: str = ""
    phone: str = ""
    email: str = ""
    website: str = ""


class ClientInfo(Flexible):
    name: str = ""
    address: str = ""
    attention: str = ""
    reference: str = ""


class OpinionMeta(Flexible):
    number: str = ""
    title: str = "LEGAL OPINION"
    subject: str = ""
    date: str = ""
    place: str = "Jakarta"
    confidentiality: str = "RAHASIA"
    prepared_by: str = ""
    reviewed_by: str = ""


class Authority(Flexible):
    citation: str = ""
    summary: str = ""
    status: str = "Berlaku"
    relevance: str = ""


class Jurisprudence(Flexible):
    citation: str = ""
    part: str = ""
    classification: str = ""
    relevance: str = ""


class Issue(Flexible):
    question: str = ""
    short_answer: str = ""


class AnalysisSection(Flexible):
    heading: str = ""
    paragraphs: list[str] = []


class Recommendation(Flexible):
    text: str = ""
    priority: str = ""
    deadline: str = ""
    owner: str = ""


class Conclusion(Flexible):
    summary: str = ""
    recommendations: list[Recommendation] = []
    risk: str = ""
    qualifications: list[str] = []
    assumptions: list[str] = []
    limitations: list[str] = []


class DocumentReview(Flexible):
    name: str = ""
    reference: str = ""
    status: str = ""


class Signature(Flexible):
    place: str = ""
    date: str = ""
    name: str = ""
    position: str = "Partner"
    firm: str = ""
    bar_number: str = ""


class Opinion(BaseModel):
    firm: FirmInfo = FirmInfo()
    client: ClientInfo = ClientInfo()
    meta: OpinionMeta = OpinionMeta()
    facts: list[str] = []
    documents: list[DocumentReview] = []
    issues: list[Issue] = []
    legal_basis: list[Authority] = []
    jurisprudence: list[Jurisprudence] = []
    framework_notes: list[str] = []
    analysis: list[AnalysisSection] = []
    counterarguments: list[str] = []
    conclusion: Conclusion = Conclusion()
    signature: Signature = Signature()
    scope: str = ""
    background: list[str] = []
    warnings: list[str] = []


# ---------------------------------------------------------------------------
# Normalisation
# ---------------------------------------------------------------------------


def _text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text or default


def _pick(data: Mapping[str, Any], *names: str, default: Any = None) -> Any:
    for name in names:
        if name in data and data[name] not in (None, ""):
            return data[name]
    return default


def _dict(value: Any) -> dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}


def _list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    return [value]


def _paragraphs(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return []
        parts = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        return parts or [text]
    return [_text(item) for item in _list(value) if _text(item)]


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"(?i)\b(?:rp|idr)\b", "", str(value)).replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    elif text.count(".") > 1:
        text = text.replace(".", "")
    try:
        return float(text)
    except ValueError:
        match = re.search(r"\d+(?:[.,]\d+)?", text)
        if not match:
            return None
        try:
            return float(match.group(0).replace(",", "."))
        except ValueError:
            return None


def _indonesian_date(value: str | None = None) -> str:
    if value:
        text = value.strip()
        if "," in text:
            return text
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return text
        try:
            parsed = datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            return text
    else:
        parsed = date.today()
    return f"{parsed.day} {MONTHS[parsed.month - 1]} {parsed.year}"


def _date_line(place: str, value: str) -> str:
    if value and "," in value:
        return value
    return f"{place}, {_indonesian_date(value)}" if place else _indonesian_date(value)


def _recommendations(value: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for item in _list(value):
        if isinstance(item, Mapping):
            result.append(
                {
                    "text": _text(_pick(item, "text", "uraian", "recommendation", "rekomendasi", "action", "description")),
                    "priority": _text(_pick(item, "priority", "prioritas", "severity", "tingkat")),
                    "deadline": _text(_pick(item, "deadline", "tenggat", "due", "timeline")),
                    "owner": _text(_pick(item, "owner", "pic", "penanggung_jawab")),
                }
            )
        else:
            result.append({"text": _text(item), "priority": "", "deadline": "", "owner": ""})
    return result


def _sections(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, Mapping):
        items: list[Any] = []
        for heading, body in value.items():
            items.append({"heading": heading, "body": body})
        value = items
    result: list[dict[str, Any]] = []
    for item in _list(value):
        if isinstance(item, Mapping):
            heading = _text(_pick(item, "heading", "judul", "title", "section"))
            paragraphs = _paragraphs(_pick(item, "paragraphs", "paragraf", "body", "uraian", "analysis", "text", "content"))
            if not paragraphs:
                single = _text(_pick(item, "description", "argument", "summary"))
                if single:
                    paragraphs = [single]
            result.append({"heading": heading, "paragraphs": paragraphs})
        else:
            result.append({"heading": "", "paragraphs": [_text(item)]})
    return [item for item in result if item["paragraphs"] or item["heading"]]


def _normalise_opinion(raw: Any) -> Opinion:
    if not isinstance(raw, Mapping):
        raise OpinionError("Opinion JSON must be an object.")
    data = dict(raw)
    for wrapper in ("opinion", "legal_opinion", "data", "payload"):
        inner = data.get(wrapper)
        if isinstance(inner, Mapping) and len(data) <= 3:
            data = {**inner, **{k: v for k, v in data.items() if k != wrapper}}

    firm_raw = _pick(data, "firm", "law_firm", "kantor", "kantor_hukum")
    if isinstance(firm_raw, str):
        firm_raw = {"name": firm_raw}
    firm = _dict(firm_raw)
    client_raw = _pick(data, "client", "klien", "addressee", "penerima")
    if isinstance(client_raw, str):
        client_raw = {"name": client_raw}
    client = _dict(client_raw)

    meta_raw = _dict(_pick(data, "document", "meta", "metadata", "opinion_meta"))
    meta = {
        "number": _text(_pick(meta_raw, "number", "nomor", "no"), _text(_pick(data, "number", "nomor"))),
        "title": _text(_pick(meta_raw, "title", "judul"), _text(_pick(data, "title", "judul"), "LEGAL OPINION")),
        "subject": _text(_pick(meta_raw, "subject", "perihal", "matter"), _text(_pick(data, "subject", "perihal", "matter"))),
        "date": _text(_pick(meta_raw, "date", "tanggal"), _text(_pick(data, "date", "tanggal"))),
        "place": _text(_pick(meta_raw, "place", "tempat", "kota"), _text(_pick(data, "place", "tempat"), "Jakarta")),
        "confidentiality": _text(_pick(meta_raw, "confidentiality", "kerahasiaan"), "RAHASIA"),
        "prepared_by": _text(_pick(meta_raw, "prepared_by", "disusun_oleh", "author")),
        "reviewed_by": _text(_pick(meta_raw, "reviewed_by", "direview_oleh", "reviewer")),
    }

    facts = _paragraphs(_pick(data, "facts", "fakta", "statement_of_facts", "fakta_hukum", "latar_belakang"))
    if not facts:
        facts = _paragraphs(_pick(data, "background", "latar_belakang"))

    documents = []
    for item in _list(_pick(data, "documents", "dokumen", "reviewed_documents")):
        if isinstance(item, Mapping):
            documents.append(
                {
                    "name": _text(_pick(item, "name", "nama", "document", "dokumen")),
                    "reference": _text(_pick(item, "reference", "referensi", "nomor", "date", "tanggal")),
                    "status": _text(_pick(item, "status", "keadaan"), "Ditinjau"),
                }
            )
        else:
            documents.append({"name": _text(item), "reference": "", "status": "Ditinjau"})

    issues = []
    for item in _list(_pick(data, "issues", "isu_hukum", "legal_issues", "permasalahan", "questions")):
        if isinstance(item, Mapping):
            issues.append(
                {
                    "question": _text(_pick(item, "question", "pertanyaan", "issue", "isu", "title")),
                    "short_answer": _text(_pick(item, "short_answer", "jawaban_singkat", "answer", "summary")),
                }
            )
        else:
            issues.append({"question": _text(item), "short_answer": ""})

    legal_basis = []
    for item in _list(_pick(data, "legal_basis", "dasar_hukum", "authorities", "regulatory_framework", "regulations")):
        if isinstance(item, Mapping):
            legal_basis.append(
                {
                    "citation": _text(_pick(item, "citation", "sitasi", "pasal", "regulation", "law", "authority")),
                    "summary": _text(_pick(item, "summary", "ringkasan", "description", "uraian", "content")),
                    "status": _text(_pick(item, "status", "keberlakuan", "validity"), "Berlaku"),
                    "relevance": _text(_pick(item, "relevance", "relevansi", "application")),
                }
            )
        else:
            legal_basis.append({"citation": _text(item), "summary": "", "status": "Berlaku", "relevance": ""})

    jurisprudence = []
    for item in _list(_pick(data, "jurisprudence", "yurisprudensi", "precedents", "putusan")):
        if isinstance(item, Mapping):
            jurisprudence.append(
                {
                    "citation": _text(_pick(item, "citation", "sitasi", "case", "putusan", "number", "nomor")),
                    "part": _text(_pick(item, "part", "bagian", "pertimbangan", "excerpt")),
                    "classification": _text(_pick(item, "classification", "kedudukan", "ratio_or_obiter", "type")),
                    "relevance": _text(_pick(item, "relevance", "relevansi", "application")),
                }
            )
        else:
            jurisprudence.append({"citation": _text(item), "part": "", "classification": "", "relevance": ""})

    analysis = _sections(_pick(data, "analysis", "analisis", "analisis_yuridis", "discussion", "legal_analysis"))
    counterarguments = _paragraphs(_pick(data, "counterarguments", "argumen_lawan", "adversarial", "devils_advocate"))

    conclusion_raw = _pick(data, "conclusion", "kesimpulan", "conclusions")
    if isinstance(conclusion_raw, str):
        conclusion_raw = {"summary": conclusion_raw}
    conclusion = _dict(conclusion_raw)
    recommendations_raw = _pick(data, "recommendations", "rekomendasi", "mitigation", "mitigasi")
    if recommendations_raw is None:
        recommendations_raw = _pick(conclusion, "recommendations", "rekomendasi", "mitigation", "mitigasi")
    conclusion = {
        "summary": _text(_pick(conclusion, "summary", "ringkasan", "kesimpulan", "conclusion")),
        "recommendations": _recommendations(recommendations_raw),
        "risk": _text(_pick(conclusion, "risk", "risiko", "risk_level", "tingkat_risiko"), _text(_pick(data, "risk", "risiko"))),
        "qualifications": _paragraphs(_pick(conclusion, "qualifications", "kualifikasi", "caveats", "catatan")),
        "assumptions": _paragraphs(_pick(conclusion, "assumptions", "asumsi")),
        "limitations": _paragraphs(_pick(conclusion, "limitations", "keterbatasan", "batasan")),
    }

    signature_raw = _dict(_pick(data, "signature", "tanda_tangan", "signer", "partner", "penandatangan"))
    signature = {
        "place": _text(_pick(signature_raw, "place", "tempat", "kota"), meta["place"]),
        "date": _text(_pick(signature_raw, "date", "tanggal"), meta["date"]),
        "name": _text(_pick(signature_raw, "name", "nama", "partner_name")),
        "position": _text(_pick(signature_raw, "position", "jabatan", "title"), "Partner"),
        "firm": _text(_pick(signature_raw, "firm", "kantor", "law_firm"), _text(firm, "name")),
        "bar_number": _text(_pick(signature_raw, "bar_number", "nomor_peradi", "peradi", "license")),
    }

    opinion = Opinion(
        firm=FirmInfo.model_validate(
            {
                "name": _text(_pick(firm, "name", "nama", "firm_name")),
                "address": _text(_pick(firm, "address", "alamat")),
                "phone": _text(_pick(firm, "phone", "telepon", "telp")),
                "email": _text(_pick(firm, "email", "surel")),
                "website": _text(_pick(firm, "website", "web")),
            }
        ),
        client=ClientInfo.model_validate(
            {
                "name": _text(_pick(client, "name", "nama", "client_name")),
                "address": _text(_pick(client, "address", "alamat")),
                "attention": _text(_pick(client, "attention", "up", "kepada")),
                "reference": _text(_pick(client, "reference", "referensi", "ref")),
            }
        ),
        meta=OpinionMeta.model_validate(meta),
        facts=facts,
        documents=[DocumentReview.model_validate(item) for item in documents],
        issues=[Issue.model_validate(item) for item in issues],
        legal_basis=[Authority.model_validate(item) for item in legal_basis],
        jurisprudence=[Jurisprudence.model_validate(item) for item in jurisprudence],
        framework_notes=_paragraphs(_pick(data, "framework_notes", "catatan_kerangka_hukum", "regulatory_notes")),
        analysis=[AnalysisSection.model_validate(item) for item in analysis],
        counterarguments=counterarguments,
        conclusion=Conclusion.model_validate(conclusion),
        signature=Signature.model_validate(signature),
        scope=_text(_pick(data, "scope", "lingkup", "ruang_lingkup")),
        background=_paragraphs(_pick(data, "background", "latar_belakang")),
        warnings=[],
    )
    return opinion


# ---------------------------------------------------------------------------
# DOCX helpers
# ---------------------------------------------------------------------------


def _set_run_font(run: Any, *, size: float = 11.5, bold: bool = False, italic: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_field(paragraph: Any, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    _set_run_font(run, size=8, color=GREY)


def _setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("RAHASIA — LEGAL OPINION")
    _set_run_font(run, size=8, bold=True, color=GREY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Halaman ")
    _set_run_font(run, size=8, color=GREY)
    _add_field(footer, "PAGE")
    run = footer.add_run(" dari ")
    _set_run_font(run, size=8, color=GREY)
    _add_field(footer, "NUMPAGES")
    return doc


def _paragraph(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    size: float = 11.5,
    align: int = WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_after: float = 6,
    space_before: float = 0,
    indent: float | None = None,
    first_line: float | None = None,
    color: RGBColor | None = None,
    keep_with_next: bool = False,
) -> Any:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.line_spacing = 1.3
    if indent is not None:
        paragraph.paragraph_format.left_indent = Cm(indent)
    if first_line is not None:
        paragraph.paragraph_format.first_line_indent = Cm(first_line)
    if keep_with_next:
        paragraph.paragraph_format.keep_with_next = True
    if text:
        run = paragraph.add_run(text)
        _set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def _numbered(doc: Document, number: str, text: str, *, indent: float = 0.9) -> Any:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.left_indent = Cm(indent)
    paragraph.paragraph_format.first_line_indent = Cm(-indent)
    run = paragraph.add_run(f"{number}. ")
    _set_run_font(run)
    run = paragraph.add_run(text)
    _set_run_font(run)
    return paragraph


def _heading(doc: Document, text: str, *, size: float = 12.5, space_before: float = 16) -> Any:
    return _paragraph(
        doc,
        text.upper(),
        bold=True,
        size=size,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=space_before,
        space_after=6,
        color=NAVY,
        keep_with_next=True,
    )


def _subheading(doc: Document, text: str) -> Any:
    return _paragraph(doc, text, bold=True, size=11.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=3, color=NAVY, keep_with_next=True)


def _cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_margins(cell: Any, top: int = 80, bottom: int = 80, left: int = 100, right: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _table_header_repeat(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cell_text(cell: Any, text: str, *, bold: bool = False, size: float = 9, color: RGBColor | None = None, align: int = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float]) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    _table_header_repeat(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        _set_cell_text(cell, label, bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_shading(cell, NAVY_FILL)
        _cell_margins(cell)
        cell.width = Cm(widths[index])
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            _set_cell_text(cell, _text(value), size=8.5)
            _cell_margins(cell)
            cell.width = Cm(widths[index])
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    return table


def _signature_block(doc: Document, signature: Signature) -> None:
    _paragraph(doc, "Hormat kami,", space_before=20, space_after=0)
    if signature.firm:
        _paragraph(doc, signature.firm, bold=True, space_after=0)
    _paragraph(doc, "", space_after=0)
    _paragraph(doc, "", space_after=0)
    _paragraph(doc, "", space_after=0)
    _paragraph(doc, signature.name or "________________________", bold=True, space_after=0)
    _paragraph(doc, signature.position or "Partner", space_after=0)
    if signature.bar_number:
        _paragraph(doc, f"Nomor Peradi: {signature.bar_number}", size=10, space_after=0)
    _paragraph(doc, _date_line(signature.place, signature.date), size=10, color=GREY, space_after=0)


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------


def build_opinion(opinion: Opinion) -> tuple[Document, dict[str, Any]]:
    doc = _setup_document()
    warnings: list[str] = []
    if not opinion.client.name:
        warnings.append("Nama klien belum diisi.")
    if not opinion.facts:
        warnings.append("Fakta hukum belum diisi; Bagian I akan kosong.")
    if not opinion.issues:
        warnings.append("Isu hukum belum diisi; Bagian II akan kosong.")
    if not opinion.legal_basis:
        warnings.append("Dasar hukum belum diisi; Bagian III akan kosong. Jangan mengarang sitasi.")
    if not opinion.conclusion.summary:
        warnings.append("Kesimpulan belum diisi; Bagian V akan kosong.")
    if not opinion.signature.name:
        warnings.append("Nama penandatangan belum diisi; dokumen memakai garis tanda tangan kosong.")

    # Cover -------------------------------------------------------------------
    if opinion.firm.name:
        _paragraph(doc, opinion.firm.name.upper(), bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color=NAVY)
    else:
        _paragraph(doc, "[NAMA KANTOR ADVOKAT]", bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color=NAVY)
    contact = " | ".join(value for value in (opinion.firm.address, opinion.firm.phone, opinion.firm.email, opinion.firm.website) if value)
    if contact:
        _paragraph(doc, contact, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=6)

    line = _paragraph(doc, "", space_after=40)
    p_pr = line._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3864")
    borders.append(bottom)
    p_pr.append(borders)

    _paragraph(doc, opinion.meta.confidentiality.upper(), bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x9C, 0x00, 0x06), space_before=20, space_after=30)
    _paragraph(doc, opinion.meta.title.upper(), bold=True, size=28, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, space_after=10)
    if opinion.meta.subject:
        _paragraph(doc, opinion.meta.subject, bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    if opinion.meta.number:
        _paragraph(doc, f"Nomor: {opinion.meta.number}", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _paragraph(doc, _date_line(opinion.meta.place, opinion.meta.date), size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)

    _paragraph(doc, "Disiapkan untuk:", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=2)
    _paragraph(doc, opinion.client.name or "[NAMA KLIEN]", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if opinion.client.attention:
        _paragraph(doc, f"u.p. {opinion.client.attention}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if opinion.client.address:
        _paragraph(doc, opinion.client.address, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=20)
    if opinion.client.reference:
        _paragraph(doc, f"Referensi: {opinion.client.reference}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=8)

    _paragraph(
        doc,
        "Dokumen ini bersifat rahasia dan disiapkan hanya untuk kepentingan klien. Dokumen ini tidak boleh diperbanyak atau diserahkan kepada pihak lain tanpa persetujuan tertulis.",
        size=9,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=GREY,
        space_before=40,
    )
    doc.add_page_break()

    # Section I ---------------------------------------------------------------
    _heading(doc, "I. FAKTA-FAKTA HUKUM TERVERIFIKASI (STATEMENT OF FACTS)", space_before=0)
    if opinion.scope:
        _paragraph(doc, f"Ruang lingkup: {opinion.scope}", italic=True, color=GREY)
    if opinion.background and not opinion.facts:
        facts = opinion.background
    else:
        facts = opinion.facts
    if facts:
        for index, fact in enumerate(facts, start=1):
            _numbered(doc, str(index), fact)
    else:
        _paragraph(doc, "Belum ada fakta hukum yang diberikan.", italic=True, color=GREY)

    if opinion.documents:
        _subheading(doc, "Dokumen yang Ditinjau")
        rows = [
            [index, item.name, item.reference, item.status]
            for index, item in enumerate(opinion.documents, start=1)
        ]
        _table(doc, ["No", "Dokumen", "Tanggal / Referensi", "Status"], rows, [1.2, 7.5, 4.0, 3.0])

    # Section II --------------------------------------------------------------
    _heading(doc, "II. RUMUSAN MASALAH YURIDIS (LEGAL ISSUES PRESENTED)")
    if opinion.issues:
        for index, issue in enumerate(opinion.issues, start=1):
            _numbered(doc, str(index), issue.question)
            if issue.short_answer:
                _paragraph(doc, f"Jawaban singkat: {issue.short_answer}", indent=0.9, italic=True)
    else:
        _paragraph(doc, "Belum ada isu hukum yang dirumuskan.", italic=True, color=GREY)

    # Section III -------------------------------------------------------------
    _heading(doc, "III. KERANGKA HUKUM & YURISPRUDENSI TERKAIT (REGULATORY & PRECEDENT FRAMEWORK)")
    if opinion.legal_basis:
        rows = [
            [
                index,
                authority.citation,
                authority.summary,
                authority.status,
                authority.relevance,
            ]
            for index, authority in enumerate(opinion.legal_basis, start=1)
        ]
        _table(doc, ["No", "Dasar Hukum", "Ringkasan Ketentuan", "Status", "Relevansi"], rows, [1.0, 4.0, 6.5, 2.0, 3.0])
    else:
        _paragraph(doc, "Belum ada dasar hukum yang diberikan. Jangan mengarang nomor peraturan atau pasal.", italic=True, color=GREY)

    if opinion.jurisprudence:
        _subheading(doc, "Yurisprudensi")
        rows = [
            [
                index,
                item.citation,
                item.part,
                item.classification,
                item.relevance,
            ]
            for index, item in enumerate(opinion.jurisprudence, start=1)
        ]
        _table(doc, ["No", "Putusan", "Bagian yang Dikutip", "Ratio / Obiter", "Relevansi"], rows, [1.0, 4.0, 5.0, 2.5, 3.0])
        _paragraph(
            doc,
            "Ratio decidendi adalah pertimbangan yang menjadi dasar amar dan mengikat pada pokok yang diputus; obiter dicta tidak mengikat. Klasifikasi di atas harus diverifikasi terhadap teks putusan.",
            size=9,
            italic=True,
            color=GREY,
        )

    for note in opinion.framework_notes:
        _paragraph(doc, note)

    # Section IV --------------------------------------------------------------
    _heading(doc, "IV. ANALISIS YURIDIS MENDALAM (SUBSTANTIVE LEGAL ANALYSIS)")
    if opinion.analysis:
        for index, section in enumerate(opinion.analysis, start=1):
            heading = section.heading or f"Analisis {index}"
            _subheading(doc, f"{chr(64 + index)}. {heading}" if index <= 26 else heading)
            for paragraph in section.paragraphs:
                _paragraph(doc, paragraph)
    else:
        _paragraph(doc, "Belum ada analisis yuridis yang diberikan.", italic=True, color=GREY)

    if opinion.counterarguments:
        _subheading(doc, "Argumen yang Bertentangan dan Uji Ketahanan Dalil")
        for index, argument in enumerate(opinion.counterarguments, start=1):
            _numbered(doc, str(index), argument)

    # Section V ---------------------------------------------------------------
    _heading(doc, "V. KESIMPULAN & REKOMENDASI LANGKAH MITIGASI (CONCLUSIONS & ACTIONABLE ADVICE)")
    if opinion.conclusion.summary:
        _paragraph(doc, opinion.conclusion.summary)
    else:
        _paragraph(doc, "Belum ada kesimpulan yang diberikan.", italic=True, color=GREY)
    if opinion.conclusion.risk:
        _paragraph(doc, f"Tingkat risiko: {opinion.conclusion.risk}.", bold=True)

    if opinion.conclusion.recommendations:
        _subheading(doc, "Rekomendasi dan Langkah Mitigasi")
        rows = [
            [
                index,
                item.text,
                item.priority,
                item.deadline,
                item.owner,
            ]
            for index, item in enumerate(opinion.conclusion.recommendations, start=1)
        ]
        _table(doc, ["No", "Rekomendasi / Tindakan", "Prioritas", "Tenggat", "PIC"], rows, [1.0, 8.0, 2.0, 2.5, 3.0])

    if opinion.conclusion.assumptions:
        _subheading(doc, "Asumsi")
        for index, assumption in enumerate(opinion.conclusion.assumptions, start=1):
            _numbered(doc, str(index), assumption)
    if opinion.conclusion.qualifications:
        _subheading(doc, "Kualifikasi dan Peringatan")
        for index, qualification in enumerate(opinion.conclusion.qualifications, start=1):
            _numbered(doc, str(index), qualification)
    if opinion.conclusion.limitations:
        _subheading(doc, "Keterbatasan")
        for index, limitation in enumerate(opinion.conclusion.limitations, start=1):
            _numbered(doc, str(index), limitation)

    _paragraph(
        doc,
        "Opini ini disusun berdasarkan fakta dan dokumen yang diberikan kepada kami. Kami tidak melakukan verifikasi independen atas kebenaran fakta, keaslian dokumen, atau status hukum pihak ketiga, kecuali dinyatakan lain. Opini ini bukan jaminan hasil dan harus ditinjau oleh advokat berlisensi sebelum digunakan untuk mengambil keputusan.",
        size=9,
        italic=True,
        color=GREY,
        space_before=12,
    )
    _signature_block(doc, opinion.signature)

    doc.core_properties.title = opinion.meta.title
    doc.core_properties.subject = opinion.meta.subject
    doc.core_properties.author = opinion.signature.firm or opinion.firm.name
    doc.core_properties.keywords = "legal opinion, Indonesia"

    report = {
        "ok": True,
        "documentType": "Legal Opinion",
        "sections": [
            "I. Fakta-Fakta Hukum Terverifikasi",
            "II. Rumusan Masalah Yuridis",
            "III. Kerangka Hukum & Yurisprudensi Terkait",
            "IV. Analisis Yuridis Mendalam",
            "V. Kesimpulan & Rekomendasi Langkah Mitigasi",
        ],
        "facts": len(facts),
        "issues": len(opinion.issues),
        "authorities": len(opinion.legal_basis),
        "jurisprudence": len(opinion.jurisprudence),
        "analysisSections": len(opinion.analysis),
        "recommendations": len(opinion.conclusion.recommendations),
        "warnings": warnings,
    }
    return doc, report


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _load_input(input_value: str) -> Opinion:
    if input_value == "-":
        try:
            raw = json.load(sys.stdin)
        except json.JSONDecodeError as exc:
            raise OpinionError(f"Input JSON dari stdin tidak valid: {exc}") from exc
    else:
        path = Path(input_value)
        if not path.exists():
            raise OpinionError(f"Input tidak ditemukan: {path}")
        if path.suffix.lower() != ".json":
            raise OpinionError("Input opini harus berupa file .json.")
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise OpinionError(f"Input JSON tidak dapat dibaca: {exc}") from exc
    return _normalise_opinion(raw)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build a formal five-part legal opinion DOCX."
    )
    parser.add_argument("--input", required=True, help="Opinion JSON path, or - for stdin.")
    parser.add_argument("--output", required=True, help="Output .docx path.")
    parser.add_argument("--json", action="store_true", help="Print a machine-readable report.")
    args = parser.parse_args(argv)

    try:
        opinion = _load_input(args.input)
        output = Path(args.output)
        if output.suffix.lower() != ".docx":
            raise OpinionError("Output harus berupa file .docx.")
        doc, report = build_opinion(opinion)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
        report["output"] = str(output)
    except (OpinionError, ValidationError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(report, ensure_ascii=False))
    else:
        print(f"Legal opinion written to {report['output']}")
        for warning in report.get("warnings", []):
            print(f"Warning: {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
