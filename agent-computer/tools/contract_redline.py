#!/usr/bin/env python3
"""Review a DOCX contract against a playbook and emit tracked changes plus comments.

The output is a new DOCX package whose revisions use the native
``w:ins``/``w:del`` elements and whose notes use the native ``w:comment``
part, so Microsoft Word shows them as Track Changes and Review comments.
"""

from __future__ import annotations

import argparse
import json
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
COMMENTS_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)
COMMENTS_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
)

MAX_REPLACEMENTS_PER_PARAGRAPH = 100
MAX_REPLACEMENTS_PER_RULE = 1000


class RedlineError(Exception):
    """A playbook or DOCX problem that should be reported to the caller."""


@dataclass
class Rule:
    id: str
    kind: str
    title: str
    reason: str
    severity: str
    patterns: list[re.Pattern[str]] = field(default_factory=list)
    keywords: list[str] = field(default_factory=list)
    replacement: str | None = None
    suggested_text: str | None = None
    case_sensitive: bool = False
    regex: bool = False


@dataclass
class CommentRecord:
    id: int
    text: str
    author: str
    date: str


@dataclass
class ChangeResult:
    okay: bool
    excerpt: str = ""
    reason: str = ""


def _as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def _first_text(data: dict[str, Any], names: Sequence[str], default: str = "") -> str:
    for name in names:
        value = data.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return default


def _first_bool(data: dict[str, Any], names: Sequence[str], default: bool) -> bool:
    for name in names:
        if name in data:
            return bool(data[name])
    return default


def _normalise_severity(value: Any, fallback: str) -> str:
    text = str(value or fallback).strip().lower()
    return {
        "tinggi": "high",
        "sedang": "medium",
        "menengah": "medium",
        "rendah": "low",
    }.get(text, text or fallback)


def _normalise_kind(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    if text in {"required", "mandatory", "wajib", "klausul_wajib"}:
        return "required"
    if text in {
        "forbidden",
        "risk",
        "risky",
        "prohibited",
        "terlarang",
        "berisiko",
        "klausul_terlarang",
    }:
        return "forbidden"
    if text in {
        "replacement",
        "replace",
        "text_replacement",
        "pergantian",
        "standar_pergantian",
    }:
        return "replacement"
    return None


def _collect_rules(data: dict[str, Any], names: Sequence[str]) -> list[dict[str, Any]]:
    collected: list[dict[str, Any]] = []
    for name in names:
        value = data.get(name)
        if value is None:
            continue
        if isinstance(value, dict):
            collected.append(value)
        elif isinstance(value, list):
            collected.extend(item for item in value if isinstance(item, dict))
        else:
            raise RedlineError(f"Playbook field {name!r} must be an object or a list.")
    return collected


def _compile_pattern(source: str, flags: int, where: str) -> re.Pattern[str]:
    try:
        return re.compile(source, flags)
    except re.error as exc:
        raise RedlineError(f"Invalid regular expression in {where}: {exc}") from exc


def _build_rule(data: dict[str, Any], kind: str, index: int) -> Rule:
    rule_id = _first_text(data, ("id", "rule_id", "slug"), f"{kind}-{index + 1}")
    title = _first_text(data, ("title", "name", "judul"), rule_id)
    reason = _first_text(
        data,
        ("reason", "comment", "note", "description", "message", "legal_reason"),
    )
    severity = _normalise_severity(
        data.get("severity")
        or data.get("risk")
        or data.get("risk_level")
        or data.get("tingkat_risiko"),
        "high" if kind in {"required", "forbidden"} else "medium",
    )
    case_sensitive = _first_bool(
        data, ("case_sensitive", "caseSensitive"), default=(kind == "replacement")
    )
    flags = 0 if case_sensitive else re.IGNORECASE

    if kind == "replacement":
        source = _first_text(
            data,
            ("find", "search", "from", "match", "text", "pattern"),
        )
        if not source:
            raise RedlineError(f"Replacement rule {rule_id!r} has no find text.")
        use_regex = _first_bool(
            data, ("regex", "use_regex", "is_regex", "regexp"), default=False
        )
        pattern_source = source if use_regex else re.escape(source)
        pattern = _compile_pattern(pattern_source, flags, f"rule {rule_id!r}")
        replacement = data.get("replace")
        if replacement is None:
            replacement = data.get("replacement")
        if replacement is None:
            replacement = data.get("to")
        if replacement is None:
            replacement = data.get("with")
        if replacement is None:
            replacement = data.get("suggested_replacement")
        if replacement is None:
            replacement = ""
        return Rule(
            id=rule_id,
            kind=kind,
            title=title,
            reason=reason,
            severity=severity,
            patterns=[pattern],
            replacement=str(replacement),
            case_sensitive=case_sensitive,
            regex=use_regex,
        )

    patterns_raw: list[str] = []
    for name in ("patterns", "pattern", "regex", "regexes"):
        for value in _as_list(data.get(name)):
            text = str(value).strip()
            if text:
                patterns_raw.append(text)

    keywords: list[str] = []
    for name in ("keywords", "keyword", "contains", "contains_any", "match"):
        for value in _as_list(data.get(name)):
            text = str(value).strip()
            if text:
                keywords.append(text)

    if not patterns_raw and not keywords:
        raise RedlineError(
            f"{kind.title()} rule {rule_id!r} needs patterns, keywords, or match text."
        )

    replacement = None
    if kind == "forbidden":
        replacement = _first_text(
            data,
            (
                "replace",
                "replacement",
                "suggested_replacement",
                "mitigation_text",
                "standard_text",
            ),
            "",
        )
        if not replacement:
            replacement = None

    suggested_text = None
    if kind == "required":
        suggested_text = _first_text(
            data,
            (
                "suggested_text",
                "insert_text",
                "standard_text",
                "suggested_clause",
                "clause_text",
            ),
            "",
        )
        if not suggested_text:
            suggested_text = None

    return Rule(
        id=rule_id,
        kind=kind,
        title=title,
        reason=reason,
        severity=severity,
        patterns=[
            _compile_pattern(source, flags, f"rule {rule_id!r}")
            for source in patterns_raw
        ],
        keywords=keywords,
        replacement=replacement,
        suggested_text=suggested_text,
        case_sensitive=case_sensitive,
    )


def load_playbook(path: Path) -> tuple[list[Rule], list[Rule], list[Rule]]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise RedlineError(f"Could not read playbook {path}: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise RedlineError(f"Playbook {path} is not valid JSON: {exc}") from exc

    if not isinstance(data, dict):
        raise RedlineError("The playbook must be a JSON object.")

    nested = data.get("rules")
    if isinstance(nested, dict):
        merged = dict(data)
        for key, value in nested.items():
            merged.setdefault(key, value)
        data = merged
    elif isinstance(nested, list):
        grouped: dict[str, list[dict[str, Any]]] = {
            "required": [],
            "forbidden": [],
            "replacement": [],
        }
        for item in nested:
            if not isinstance(item, dict):
                continue
            kind = _normalise_kind(item.get("type") or item.get("kind"))
            if kind:
                grouped[kind].append(item)
        merged = dict(data)
        data = merged
        for kind, items in grouped.items():
            if not items:
                continue
            key = {
                "required": "required_clauses",
                "forbidden": "forbidden_clauses",
                "replacement": "replacements",
            }[kind]
            existing = data.get(key)
            if existing is None:
                data[key] = items
            elif isinstance(existing, list):
                data[key] = items + existing

    required = _collect_rules(
        data, ("required_clauses", "mandatory_clauses", "klausul_wajib", "required")
    )
    forbidden = _collect_rules(
        data,
        (
            "forbidden_clauses",
            "risk_rules",
            "risky_clauses",
            "prohibited_clauses",
            "klausul_terlarang",
            "klausul_berisiko",
        ),
    )
    replacements = _collect_rules(
        data,
        (
            "replacements",
            "text_replacements",
            "standard_replacements",
            "standar_pergantian",
        ),
    )

    if not required and not forbidden and not replacements:
        raise RedlineError("The playbook contains no recognised rules.")

    return (
        [_build_rule(item, "required", i) for i, item in enumerate(required)],
        [_build_rule(item, "forbidden", i) for i, item in enumerate(forbidden)],
        [_build_rule(item, "replacement", i) for i, item in enumerate(replacements)],
    )


def _inside(element: etree._Element, tag: str) -> bool:
    parent = element.getparent()
    while parent is not None:
        if parent.tag == qn(tag):
            return True
        parent = parent.getparent()
    return False


def paragraph_text(paragraph: etree._Element) -> str:
    parts: list[str] = []
    for node in paragraph.iter(qn("w:t")):
        if not _inside(node, "w:del"):
            parts.append(node.text or "")
    return "".join(parts)


def direct_paragraph_text(paragraph: etree._Element) -> str:
    return "".join(
        _run_text(run) for run in paragraph if run.tag == qn("w:r")
    )


def _run_text(run: etree._Element) -> str:
    return "".join(node.text or "" for node in run.findall(qn("w:t")))


def _run_rpr(run: etree._Element) -> etree._Element | None:
    rpr = run.find(qn("w:rPr"))
    return deepcopy(rpr) if rpr is not None else None


def _text_run(
    source_run: etree._Element,
    text: str,
    *,
    deleted: bool = False,
    rpr: etree._Element | None = None,
) -> etree._Element:
    run = OxmlElement("w:r")
    properties = rpr if rpr is not None else _run_rpr(source_run)
    if properties is not None:
        run.append(properties)
    node = OxmlElement("w:delText" if deleted else "w:t")
    node.set(f"{{{XML_NS}}}space", "preserve")
    node.text = text
    run.append(node)
    return run


def _has_unsupported_children(run: etree._Element) -> bool:
    allowed = {qn("w:rPr"), qn("w:t")}
    return any(child.tag not in allowed for child in run)


def _comment_reference_run(comment_id: int) -> etree._Element:
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "CommentReference")
    rpr.append(style)
    run.append(rpr)
    reference = OxmlElement("w:commentReference")
    reference.set(qn("w:id"), str(comment_id))
    run.append(reference)
    return run


def _anchor_whole_paragraph(paragraph: etree._Element, comment_id: int) -> None:
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(comment_id))
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(comment_id))
    ppr = paragraph.find(qn("w:pPr"))
    paragraph.insert(1 if ppr is not None else 0, start)
    paragraph.append(end)
    paragraph.append(_comment_reference_run(comment_id))


def _max_revision_id(document: Document) -> int:
    highest = 0
    revision_tags = {
        qn("w:ins"),
        qn("w:del"),
        qn("w:moveFrom"),
        qn("w:moveTo"),
    }
    for element in document.element.body.iter():
        if element.tag not in revision_tags:
            continue
        try:
            highest = max(highest, int(element.get(qn("w:id"), "0")))
        except ValueError:
            continue
    return highest


def _existing_comment_ids(path: Path) -> list[int]:
    try:
        with zipfile.ZipFile(path) as package:
            candidates = [
                name
                for name in package.namelist()
                if name.startswith("word/comments") and name.endswith(".xml")
            ]
            if not candidates:
                return []
            root = etree.fromstring(package.read(candidates[0]))
    except (OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return []
    ids: list[int] = []
    for comment in root.findall(qn("w:comment")):
        try:
            ids.append(int(comment.get(qn("w:id"), "")))
        except ValueError:
            continue
    return ids


class Redliner:
    def __init__(
        self,
        document: Document,
        *,
        author: str,
        date: str,
        existing_comment_ids: Iterable[int],
        warnings: list[str],
    ) -> None:
        self.document = document
        self.author = author
        self.date = date
        self.warnings = warnings
        self._revision_id = _max_revision_id(document) + 1
        self._comment_id = max(list(existing_comment_ids) or [-1]) + 1
        self.comments: list[CommentRecord] = []
        self.findings: list[dict[str, Any]] = []
        self.paragraphs: list[etree._Element] = list(
            document.element.body.iter(qn("w:p"))
        )
        self.replacements = 0
        self.insertions = 0
        self.deletions = 0
        self.unmatched: set[str] = set()

    def apply(
        self,
        required: Sequence[Rule],
        forbidden: Sequence[Rule],
        replacements: Sequence[Rule],
    ) -> None:
        for rule in forbidden:
            self._apply_risk_rule(rule)
        for rule in replacements:
            self._apply_replacement_rule(rule)
        for rule in required:
            self._apply_required_rule(rule)

    def _new_revision_id(self) -> int:
        value = self._revision_id
        self._revision_id += 1
        return value

    def _new_comment(self, text: str) -> int:
        comment_id = self._comment_id
        self._comment_id += 1
        self.comments.append(
            CommentRecord(comment_id, text.strip(), self.author, self.date)
        )
        return comment_id

    def _matches(self, rule: Rule, text: str) -> bool:
        for pattern in rule.patterns:
            if pattern.search(text):
                return True
        haystack = text if rule.case_sensitive else text.casefold()
        for keyword in rule.keywords:
            needle = keyword if rule.case_sensitive else keyword.casefold()
            if needle in haystack:
                return True
        return False

    def _matching_excerpts(self, rule: Rule, text: str) -> list[str]:
        excerpts: list[str] = []
        for pattern in rule.patterns:
            for match in pattern.finditer(text):
                found = match.group(0).strip()
                if found and found not in excerpts:
                    excerpts.append(found)
                if len(excerpts) >= 3:
                    return excerpts
        haystack = text if rule.case_sensitive else text.casefold()
        for keyword in rule.keywords:
            needle = keyword if rule.case_sensitive else keyword.casefold()
            position = haystack.find(needle)
            if position < 0:
                continue
            start = max(0, position - 40)
            end = min(len(text), position + len(keyword) + 40)
            excerpt = " ".join(text[start:end].split())
            if excerpt and excerpt not in excerpts:
                excerpts.append(excerpt)
            if len(excerpts) >= 3:
                break
        return excerpts

    def _comment_text(
        self, rule: Rule, label: str, excerpt: str = "", note: str = ""
    ) -> str:
        lines = [f"[{rule.severity.upper()}] {label}: {rule.title}"]
        if rule.reason:
            lines.append(rule.reason)
        if note:
            lines.append(note)
        if excerpt:
            clipped = " ".join(excerpt.split())
            if len(clipped) > 300:
                clipped = clipped[:297].rstrip() + "..."
            lines.append(f'Kutipan: "{clipped}"')
        return "\n".join(lines)

    def _apply_risk_rule(self, rule: Rule) -> None:
        matched = False
        for index, paragraph in enumerate(self.paragraphs, start=1):
            text = paragraph_text(paragraph)
            if not self._matches(rule, text):
                continue
            matched = True
            if rule.replacement is not None:
                changed = self._replace_all(
                    paragraph,
                    rule,
                    index,
                    label="Klausul berisiko",
                    only_direct=True,
                )
                if changed:
                    continue
            excerpts = self._matching_excerpts(rule, text)
            excerpt = excerpts[0] if excerpts else ""
            note = f"Ditemukan {len(excerpts)} kecocokan." if len(excerpts) > 1 else ""
            comment = self._comment_text(rule, "Klausul berisiko", excerpt, note)
            comment_id = self._new_comment(comment)
            _anchor_whole_paragraph(paragraph, comment_id)
            self.findings.append(
                {
                    "rule": rule.id,
                    "type": rule.kind,
                    "severity": rule.severity,
                    "paragraph": index,
                    "action": "commented",
                    "excerpt": excerpt,
                    "comment": comment,
                }
            )
        if not matched:
            self.unmatched.add(rule.id)

    def _apply_replacement_rule(self, rule: Rule) -> None:
        matched = False
        for index, paragraph in enumerate(self.paragraphs, start=1):
            count = self._replace_all(
                paragraph,
                rule,
                index,
                label="Standar pergantian teks",
                only_direct=True,
            )
            if count:
                matched = True
        if not matched:
            self.unmatched.add(rule.id)
            self.warnings.append(
                f"Replacement rule {rule.id!r} did not match any paragraph text."
            )

    def _apply_required_rule(self, rule: Rule) -> None:
        for paragraph in self.paragraphs:
            if self._matches(rule, paragraph_text(paragraph)):
                return

        if rule.suggested_text:
            comment = self._comment_text(
                rule,
                "Klausul wajib ditambahkan",
                note="Teks disisipkan sebagai Track Changes.",
            )
            comment_id = self._new_comment(comment)
            paragraph = self._append_inserted_paragraph(rule, comment_id)
            self.paragraphs.append(paragraph)
            self.insertions += 1
            self.findings.append(
                {
                    "rule": rule.id,
                    "type": rule.kind,
                    "severity": rule.severity,
                    "paragraph": len(self.paragraphs),
                    "action": "inserted",
                    "comment": comment,
                }
            )
            return

        target = self.paragraphs[0] if self.paragraphs else self._blank_paragraph()
        comment = self._comment_text(
            rule,
            "Klausul wajib tidak ditemukan",
            note="Tinjau apakah klausul ini perlu ditambahkan sebelum penandatanganan.",
        )
        comment_id = self._new_comment(comment)
        _anchor_whole_paragraph(target, comment_id)
        self.findings.append(
            {
                "rule": rule.id,
                "type": rule.kind,
                "severity": rule.severity,
                "paragraph": 1,
                "action": "commented",
                "comment": comment,
            }
        )

    def _blank_paragraph(self) -> etree._Element:
        paragraph = OxmlElement("w:p")
        body = self.document.element.body
        sect = body.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(paragraph)
        else:
            body.append(paragraph)
        self.paragraphs.insert(0, paragraph)
        return paragraph

    def _append_inserted_paragraph(
        self, rule: Rule, comment_id: int
    ) -> etree._Element:
        paragraph = OxmlElement("w:p")
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), str(self._new_revision_id()))
        insertion.set(qn("w:author"), self.author)
        insertion.set(qn("w:date"), self.date)
        lines = (rule.suggested_text or "").split("\n")
        for offset, line in enumerate(lines):
            run = OxmlElement("w:r")
            if offset:
                run.append(OxmlElement("w:br"))
            node = OxmlElement("w:t")
            node.set(f"{{{XML_NS}}}space", "preserve")
            node.text = line
            run.append(node)
            insertion.append(run)
        paragraph.append(insertion)

        start = OxmlElement("w:commentRangeStart")
        start.set(qn("w:id"), str(comment_id))
        end = OxmlElement("w:commentRangeEnd")
        end.set(qn("w:id"), str(comment_id))
        insertion.addprevious(start)
        insertion.addnext(end)
        end.addnext(_comment_reference_run(comment_id))

        body = self.document.element.body
        sect = body.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(paragraph)
        else:
            body.append(paragraph)
        return paragraph

    def _replace_all(
        self,
        paragraph: etree._Element,
        rule: Rule,
        paragraph_index: int,
        *,
        label: str,
        only_direct: bool,
    ) -> int:
        if not rule.patterns:
            return 0
        pattern = rule.patterns[0]
        count = 0
        while count < MAX_REPLACEMENTS_PER_PARAGRAPH:
            text = (
                direct_paragraph_text(paragraph)
                if only_direct
                else paragraph_text(paragraph)
            )
            match = pattern.search(text)
            if match is None:
                break
            if match.start() == match.end():
                self.warnings.append(
                    f"Rule {rule.id!r} matched an empty string in paragraph "
                    f"{paragraph_index}; skipped."
                )
                break
            replacement = (
                match.expand(rule.replacement or "") if rule.regex else (rule.replacement or "")
            )
            result = self._replace_match(
                paragraph,
                text,
                match.start(),
                match.end(),
                replacement,
                rule,
                paragraph_index,
                label,
            )
            if not result.okay:
                if result.reason:
                    self.warnings.append(result.reason)
                break
            count += 1
            self.replacements += 1
        if count >= MAX_REPLACEMENTS_PER_PARAGRAPH:
            self.warnings.append(
                f"Rule {rule.id!r} hit the per-paragraph replacement cap in "
                f"paragraph {paragraph_index}."
            )
        return count

    def _replace_match(
        self,
        paragraph: etree._Element,
        text: str,
        start: int,
        end: int,
        replacement: str,
        rule: Rule,
        paragraph_index: int,
        label: str,
    ) -> ChangeResult:
        runs = [child for child in paragraph if child.tag == qn("w:r")]
        spans: list[tuple[etree._Element, int, int, str]] = []
        cursor = 0
        for run in runs:
            run_text = _run_text(run)
            spans.append((run, cursor, cursor + len(run_text), run_text))
            cursor += len(run_text)
        if end > cursor:
            return ChangeResult(False)

        affected = [
            span for span in spans if span[1] < end and span[2] > start
        ]
        if not affected:
            return ChangeResult(False)
        for run, _, _, _ in affected:
            if _has_unsupported_children(run):
                return ChangeResult(
                    False,
                    reason=(
                        f"Paragraph {paragraph_index}: run with tabs, breaks or "
                        f"fields cannot be redlined safely; rule {rule.id!r} skipped."
                    ),
                )

        first_run, first_start, _, first_text = affected[0]
        last_run, last_start, _, last_text = affected[-1]
        prefix = (
            first_text[: max(0, start - first_start)] if start > first_start else ""
        )
        last_offset = end - last_start
        suffix = last_text[last_offset:] if last_offset < len(last_text) else ""

        deleted_segments: list[tuple[etree._Element, str]] = []
        for run, run_start, _, run_text in affected:
            local_start = max(0, start - run_start)
            local_end = min(len(run_text), end - run_start)
            if local_start < local_end:
                deleted_segments.append((run, run_text[local_start:local_end]))

        position = paragraph.index(first_run)
        for run, _, _, _ in affected:
            paragraph.remove(run)

        nodes: list[etree._Element] = []
        if prefix:
            nodes.append(_text_run(first_run, prefix))

        deletion_node: etree._Element | None = None
        if deleted_segments:
            deletion_node = OxmlElement("w:del")
            deletion_node.set(qn("w:id"), str(self._new_revision_id()))
            deletion_node.set(qn("w:author"), self.author)
            deletion_node.set(qn("w:date"), self.date)
            for run, piece in deleted_segments:
                deletion_node.append(_text_run(run, piece, deleted=True))
            nodes.append(deletion_node)
            self.deletions += 1

        insertion_node: etree._Element | None = None
        if replacement:
            insertion_node = OxmlElement("w:ins")
            insertion_node.set(qn("w:id"), str(self._new_revision_id()))
            insertion_node.set(qn("w:author"), self.author)
            insertion_node.set(qn("w:date"), self.date)
            insertion_node.append(_text_run(first_run, replacement))
            nodes.append(insertion_node)
            self.insertions += 1

        if suffix:
            nodes.append(_text_run(last_run, suffix))

        if not nodes:
            return ChangeResult(False)
        for offset, node in enumerate(nodes):
            paragraph.insert(position + offset, node)

        excerpt = text[start:end]
        comment = self._comment_text(rule, label, excerpt)
        comment_id = self._new_comment(comment)
        first_node = (
            deletion_node
            if deletion_node is not None
            else insertion_node
            if insertion_node is not None
            else nodes[0]
        )
        last_node = (
            insertion_node
            if insertion_node is not None
            else deletion_node
            if deletion_node is not None
            else nodes[-1]
        )
        start_marker = OxmlElement("w:commentRangeStart")
        start_marker.set(qn("w:id"), str(comment_id))
        end_marker = OxmlElement("w:commentRangeEnd")
        end_marker.set(qn("w:id"), str(comment_id))
        first_node.addprevious(start_marker)
        last_node.addnext(end_marker)
        end_marker.addnext(_comment_reference_run(comment_id))

        self.findings.append(
            {
                "rule": rule.id,
                "type": rule.kind,
                "severity": rule.severity,
                "paragraph": paragraph_index,
                "action": "replaced",
                "from": excerpt,
                "to": replacement,
                "comment": comment,
            }
        )
        return ChangeResult(True, excerpt=excerpt)


def _comment_element(record: CommentRecord) -> etree._Element:
    comment = etree.Element(qn("w:comment"))
    comment.set(qn("w:id"), str(record.id))
    comment.set(qn("w:author"), record.author)
    comment.set(qn("w:date"), record.date)
    initials = "".join(word[0] for word in record.author.split()[:2]).upper()
    if initials:
        comment.set(qn("w:initials"), initials)
    for line in record.text.splitlines() or [""]:
        paragraph = etree.SubElement(comment, qn("w:p"))
        run = etree.SubElement(paragraph, qn("w:r"))
        node = etree.SubElement(run, qn("w:t"))
        node.set(f"{{{XML_NS}}}space", "preserve")
        node.text = line
    return comment


def _comments_part_and_rels(
    data: bytes | None,
) -> tuple[str, bytes]:
    target = "word/comments.xml"
    if data is None:
        root = etree.Element(f"{{{REL_NS}}}Relationships", nsmap={None: REL_NS})
    else:
        root = etree.fromstring(data)
        for relationship in root.findall(f"{{{REL_NS}}}Relationship"):
            if relationship.get("Type") != COMMENTS_REL:
                continue
            found = relationship.get("Target") or target
            if found.startswith("/"):
                target = found.lstrip("/")
            else:
                target = posixpath.normpath(posixpath.join("word", found))
            return target, etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

    used = {
        relationship.get("Id")
        for relationship in root.findall(f"{{{REL_NS}}}Relationship")
    }
    identifier = "rIdLegalComments"
    suffix = 1
    while identifier in used:
        identifier = f"rIdLegalComments{suffix}"
        suffix += 1
    relationship = etree.SubElement(root, f"{{{REL_NS}}}Relationship")
    relationship.set("Id", identifier)
    relationship.set("Type", COMMENTS_REL)
    relationship.set("Target", "comments.xml")
    return target, etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _merge_comments(
    data: bytes | None, records: Sequence[CommentRecord]
) -> bytes:
    if data is None:
        root = etree.Element(qn("w:comments"), nsmap={"w": W_NS})
    else:
        root = etree.fromstring(data)
    existing = {
        int(comment.get(qn("w:id"), "-1"))
        for comment in root.findall(qn("w:comment"))
        if (comment.get(qn("w:id")) or "").lstrip("-").isdigit()
    }
    for record in records:
        if record.id in existing:
            raise RedlineError(
                f"Comment id {record.id} already exists in the document."
            )
        root.append(_comment_element(record))
    return etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _merge_content_types(data: bytes | None, part_name: str) -> bytes:
    if data is None:
        root = etree.Element(f"{{{CT_NS}}}Types", nsmap={None: CT_NS})
    else:
        root = etree.fromstring(data)
    wanted = "/" + part_name.lstrip("/")
    for override in root.findall(f"{{{CT_NS}}}Override"):
        if override.get("PartName") == wanted:
            return etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
    override = etree.SubElement(root, f"{{{CT_NS}}}Override")
    override.set("PartName", wanted)
    override.set("ContentType", COMMENTS_CT)
    return etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _write_package(
    source_path: Path,
    output_path: Path,
    comments: Sequence[CommentRecord],
) -> None:
    if not comments:
        output_path.write_bytes(source_path.read_bytes())
        return

    rels_name = "word/_rels/document.xml.rels"
    content_types_name = "[Content_Types].xml"
    with zipfile.ZipFile(source_path) as source:
        names = set(source.namelist())
        rels_data = source.read(rels_name) if rels_name in names else None
        comments_part, rels_data = _comments_part_and_rels(rels_data)
        existing_comments = (
            source.read(comments_part) if comments_part in names else None
        )
        merged_comments = _merge_comments(existing_comments, comments)
        content_types = _merge_content_types(
            source.read(content_types_name) if content_types_name in names else None,
            comments_part,
        )

        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == rels_name:
                    data = rels_data
                elif info.filename == comments_part:
                    data = merged_comments
                elif info.filename == content_types_name:
                    data = content_types
                target.writestr(info, data)
            if rels_name not in names:
                target.writestr(rels_name, rels_data)
            if comments_part not in names:
                target.writestr(comments_part, merged_comments)
            if content_types_name not in names:
                target.writestr(content_types_name, content_types)


def _validate_package(path: Path) -> None:
    with zipfile.ZipFile(path) as package:
        broken = package.testzip()
        if broken is not None:
            raise RedlineError(f"The generated DOCX has a corrupt entry: {broken}")
        for name in ("[Content_Types].xml", "word/document.xml"):
            if name in package.namelist():
                etree.fromstring(package.read(name))
    Document(path)


def redline(
    input_path: Path,
    playbook_path: Path,
    output_path: Path,
    *,
    author: str,
    date: str,
    report_path: Path | None = None,
) -> dict[str, Any]:
    required, forbidden, replacements = load_playbook(playbook_path)
    warnings: list[str] = []
    try:
        document = Document(input_path)
    except Exception as exc:
        raise RedlineError(f"Could not open {input_path} as a DOCX file: {exc}") from exc

    redliner = Redliner(
        document,
        author=author,
        date=date,
        existing_comment_ids=_existing_comment_ids(input_path),
        warnings=warnings,
    )
    redliner.apply(required, forbidden, replacements)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".redline-", dir=output_path.parent
    ) as workspace:
        document_tmp = Path(workspace) / "document.docx"
        package_tmp = Path(workspace) / "package.docx"
        document.save(document_tmp)
        _write_package(document_tmp, package_tmp, redliner.comments)
        _validate_package(package_tmp)
        os.replace(package_tmp, output_path)

    report: dict[str, Any] = {
        "ok": True,
        "input": str(input_path),
        "playbook": str(playbook_path),
        "output": str(output_path),
        "author": author,
        "date": date,
        "summary": {
            "paragraphs": len(redliner.paragraphs),
            "replacements": redliner.replacements,
            "insertions": redliner.insertions,
            "deletions": redliner.deletions,
            "comments": len(redliner.comments),
            "requiredRules": len(required),
            "forbiddenRules": len(forbidden),
            "replacementRules": len(replacements),
        },
        "findings": redliner.findings,
        "unmatchedRules": sorted(redliner.unmatched),
        "warnings": warnings,
    }
    if report_path is not None:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    return report


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Apply playbook rules to a DOCX contract with native Track Changes and Word comments."
    )
    parser.add_argument("--input", required=True, help="Input DOCX path.")
    parser.add_argument("--playbook", required=True, help="Playbook JSON path.")
    parser.add_argument("--output", required=True, help="Output DOCX path.")
    parser.add_argument(
        "--author",
        default="Legal Agent",
        help="Track Changes and comment author name.",
    )
    parser.add_argument(
        "--date",
        default=None,
        help="ISO-8601 revision date. Defaults to the current UTC time.",
    )
    parser.add_argument("--report", default=None, help="Optional JSON report path.")
    parser.add_argument(
        "--json", action="store_true", help="Print a machine-readable report."
    )
    args = parser.parse_args(argv)

    date = args.date or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    try:
        report = redline(
            Path(args.input),
            Path(args.playbook),
            Path(args.output),
            author=args.author,
            date=date,
            report_path=Path(args.report) if args.report else None,
        )
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False),
              file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(report, ensure_ascii=False))
    else:
        summary = report["summary"]
        print(f"Redlined contract written to {report['output']}")
        print(
            f"{summary['replacements']} replacements, "
            f"{summary['comments']} comments, "
            f"{len(report['warnings'])} warnings."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
