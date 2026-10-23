#!/usr/bin/env python3
"""Build a styled five-sheet legal due diligence workbook.

The input is either a data room directory (inventoried and mined for facts) or
a JSON analysis produced by a reviewer. The output is an .xlsx workbook with
the five sheets the legal team uses for corporate, permit, financing and
litigation findings.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, ValidationError


SHEET_SUMMARY = "Executive Summary & Red Flags"
SHEET_OWNERSHIP = "Struktur Saham & Direksi"
SHEET_PERMITS = "Perizinan & OSS-RBA"
SHEET_COVENANTS = "Negative Covenants & Utang Bank"
SHEET_LITIGATION = "Litigasi & Sengketa Aset"

MAX_DATA_ROOM_FILES = 3000
MAX_TEXT_PER_FILE = 250_000
MAX_WORKBOOK_ROWS = 5000

NAVY = "1F3864"
BLUE = "2F5597"
LIGHT_BLUE = "D9E2F3"
PALE_BLUE = "EDF2FA"
GREY = "F2F2F2"
WHITE = "FFFFFF"
RED_FILL = "F4CCCC"
AMBER_FILL = "FCE5CD"
GREEN_FILL = "D9EAD3"
DARK_RED = "9C0006"
DARK_AMBER = "9C6500"
DARK_GREEN = "006100"

THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class DueDiligenceError(Exception):
    """A data room or analysis problem that should be shown to the caller."""


# ---------------------------------------------------------------------------
# Pydantic models
# ---------------------------------------------------------------------------


class Flexible(BaseModel):
    model_config = ConfigDict(extra="allow", populate_by_name=True)


class Deal(Flexible):
    name: str = ""
    client: str = ""
    target: str = ""
    reviewer: str = ""
    date: str = ""
    objective: str = ""
    source: str = ""


class Company(Flexible):
    name: str = ""
    nib: str = ""
    npwp: str = ""
    domicile: str = ""
    established: str = ""
    deed_number: str = ""
    deed_date: str = ""
    notary: str = ""
    sk_number: str = ""
    sk_date: str = ""
    business: str = ""
    notes: str = ""


class Shareholder(Flexible):
    name: str = ""
    share_class: str = ""
    shares: float | None = None
    nominal: float | None = None
    percentage: float | None = None
    since: str = ""
    source: str = ""
    notes: str = ""


class Person(Flexible):
    name: str = ""
    position: str = ""
    period: str = ""
    basis: str = ""
    status: str = ""
    source: str = ""


class Kbli(Flexible):
    code: str = ""
    title: str = ""
    risk: str = ""
    match: str = ""
    notes: str = ""
    source: str = ""


class Permit(Flexible):
    name: str = ""
    number: str = ""
    issued: str = ""
    expiry: str = ""
    issuer: str = ""
    status: str = ""
    kbli: list[Kbli] = []
    match: str = ""
    gap: str = ""
    severity: str = ""
    action: str = ""
    source: str = ""
    notes: str = ""


class Covenant(Flexible):
    category: str = ""
    document: str = ""
    creditor: str = ""
    clause: str = ""
    summary: str = ""
    restriction: str = ""
    threshold: str = ""
    actual: str = ""
    status: str = ""
    severity: str = ""
    action: str = ""
    source: str = ""
    notes: str = ""


class Loan(Flexible):
    document: str = ""
    creditor: str = ""
    facility: str = ""
    principal: float | None = None
    balance: float | None = None
    currency: str = "IDR"
    due: str = ""
    clause: str = ""
    restriction: str = ""
    status: str = ""
    severity: str = ""
    source: str = ""
    notes: str = ""


class Litigation(Flexible):
    case_number: str = ""
    court: str = ""
    parties: str = ""
    role: str = ""
    subject: str = ""
    value: float | None = None
    status: str = ""
    severity: str = ""
    asset_impact: str = ""
    source: str = ""
    action: str = ""
    notes: str = ""


class Asset(Flexible):
    name: str = ""
    asset_type: str = ""
    ownership: str = ""
    certificate: str = ""
    encumbrance: str = ""
    dispute: str = ""
    status: str = ""
    severity: str = ""
    source: str = ""
    notes: str = ""


class RedFlag(Flexible):
    severity: str = ""
    issue: str = ""
    basis: str = ""
    consequence: str = ""
    action: str = ""
    owner: str = ""
    deadline: str = ""


class DocumentInfo(Flexible):
    name: str = ""
    category: str = ""
    path: str = ""
    status: str = ""
    size: int | None = None
    modified: str = ""
    notes: str = ""


class Analysis(Flexible):
    deal: Deal = Deal()
    company: Company = Company()
    shareholders: list[Shareholder] = []
    directors: list[Person] = []
    commissioners: list[Person] = []
    corporate_actions: list[dict[str, Any]] = []
    permits: list[Permit] = []
    kbli: list[Kbli] = []
    loans: list[Loan] = []
    covenants: list[Covenant] = []
    litigation: list[Litigation] = []
    assets: list[Asset] = []
    red_flags: list[RedFlag] = []
    documents: list[DocumentInfo] = []


# ---------------------------------------------------------------------------
# Small normalisation helpers
# ---------------------------------------------------------------------------


def _text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text or default


def _pick(data: Mapping[str, Any], *names: str, default: Any = None) -> Any:
    for name in names:
        if name in data and data[name] not in (None, ""):
            return data[name]
    return default


def _dict(value: Any) -> dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}


def _list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    return [value]


def _dicts(value: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    if isinstance(value, Mapping):
        nested = _pick(value, "items", "results", "data", "rows")
        if isinstance(nested, list):
            value = nested
        else:
            return [dict(value)]
    for item in _list(value):
        if isinstance(item, Mapping):
            result.append(dict(item))
        elif item is not None:
            result.append({"name": str(item)})
    return result


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"(?i)\b(?:rp|idr|usd|sgd|eur)\b", "", text)
    text = text.replace("%", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) == 2 and len(parts[1]) in (1, 2):
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "." in text:
        parts = text.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
            text = text.replace(".", "")
    try:
        number = float(text)
    except ValueError:
        match = re.search(r"-?\d+(?:\.\d+)?", text)
        if not match:
            return None
        try:
            number = float(match.group(0))
        except ValueError:
            return None
    return -number if negative else number


def _percentage(value: Any) -> float | None:
    number = _number(value)
    if number is None:
        return None
    if 0 < number <= 1 and "%" not in str(value):
        number *= 100
    return number


def _severity(value: Any, default: str = "MEDIUM") -> str:
    text = _text(value).lower()
    if any(word in text for word in ("high", "tinggi", "kritis", "critical", "merah")):
        return "HIGH"
    if any(word in text for word in ("medium", "sedang", "menengah", "kuning")):
        return "MEDIUM"
    if any(word in text for word in ("low", "rendah", "hijau")):
        return "LOW"
    return default


def _gap_status(value: Any, default: str = "Perlu Klarifikasi") -> str:
    text = _text(value).lower()
    if any(word in text for word in ("tidak terpenuhi", "tidak memenuhi", "non-compliant", "noncompliant", "gagal", "breach", "pelanggaran")):
        return "Tidak Terpenuhi"
    if any(word in text for word in ("terpenuhi", "memenuhi", "compliant", "sesuai", "patuh", "ok", "clear")):
        return "Terpenuhi"
    if any(word in text for word in ("tidak ditemukan", "not found", "missing", "tidak ada")):
        return "Tidak Ditemukan"
    if any(word in text for word in ("klarifikasi", "clarify", "unknown", "belum", "verifikasi")):
        return "Perlu Klarifikasi"
    return default


def _permit_status(value: Any) -> str:
    text = _text(value).lower()
    if any(word in text for word in ("kedaluwarsa", "kadaluarsa", "expired", "tidak berlaku")):
        return "Kedaluwarsa"
    if any(word in text for word in ("akan berakhir", "segera berakhir", "expiring")):
        return "Akan Berakhir"
    if any(word in text for word in ("dicabut", "revoked", "dibatalkan")):
        return "Dicabut"
    if any(word in text for word in ("berlaku", "aktif", "active", "valid")):
        return "Berlaku"
    if any(word in text for word in ("tidak ditemukan", "missing", "not found")):
        return "Tidak Ditemukan"
    return "Perlu Klarifikasi"


def _litigation_status(value: Any) -> str:
    text = _text(value).lower()
    if any(word in text for word in ("selesai", "putus", "completed", "closed", "inkracht")):
        return "Selesai"
    if any(word in text for word in ("kasasi", "cassation")):
        return "Kasasi"
    if any(word in text for word in ("banding", "appeal")):
        return "Banding"
    if any(word in text for word in ("peninjauan kembali", " pk ")):
        return "Peninjauan Kembali"
    if any(word in text for word in ("berjalan", "proses", "active", "pending", "ongoing")):
        return "Berjalan"
    if any(word in text for word in ("tidak ada", "tidak ditemukan", "none", "no litigation")):
        return "Tidak Ditemukan"
    return "Perlu Klarifikasi"


def _as_paragraph_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, Mapping):
        return [f"{_text(key)}: {_text(item)}" for key, item in value.items() if _text(item)]
    return [_text(item) for item in _list(value) if _text(item)]


# ---------------------------------------------------------------------------
# JSON analysis normalisation
# ---------------------------------------------------------------------------


def _unwrap_analysis(raw: Any) -> dict[str, Any]:
    if isinstance(raw, list):
        return {"documents": raw}
    if not isinstance(raw, Mapping):
        raise DueDiligenceError("The analysis JSON must be an object or an array of documents.")
    data = dict(raw)
    for wrapper in ("analysis", "data", "result"):
        inner = data.get(wrapper)
        if isinstance(inner, Mapping):
            data = {**inner, **{k: v for k, v in data.items() if k != wrapper}}
            break
    return data


def _normalise_shareholder(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": _text(_pick(item, "name", "nama", "pemegang_saham", "shareholder", "investor")),
        "share_class": _text(_pick(item, "share_class", "class", "jenis_saham", "kategori_saham")),
        "shares": _number(_pick(item, "shares", "jumlah_saham", "lembar_saham", "quantity")),
        "nominal": _number(_pick(item, "nominal", "nilai_nominal", "nominal_per_saham", "value")),
        "percentage": _percentage(_pick(item, "percentage", "persentase", "porsi", "pct")),
        "since": _text(_pick(item, "since", "sejak", "tanggal", "date")),
        "source": _text(_pick(item, "source", "sumber", "dokumen", "document")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_person(item: dict[str, Any], position: str = "") -> dict[str, Any]:
    return {
        "name": _text(_pick(item, "name", "nama", "nama_lengkap")),
        "position": _text(_pick(item, "position", "jabatan", "role", "title"), position),
        "period": _text(_pick(item, "period", "periode", "masa_jabatan", "from", "sejak")),
        "basis": _text(_pick(item, "basis", "dasar", "dasar_hukum", "document", "dokumen", "akta")),
        "status": _text(_pick(item, "status", "keadaan"), "Perlu Klarifikasi"),
        "source": _text(_pick(item, "source", "sumber", "dokumen")),
    }


def _normalise_kbli(item: dict[str, Any] | str) -> dict[str, Any]:
    if not isinstance(item, Mapping):
        code = _text(item)
        return {"code": code, "title": "", "risk": "", "match": "Perlu Klarifikasi", "notes": "", "source": ""}
    return {
        "code": _text(_pick(item, "code", "kode", "kode_kbli", "kbli", "id")),
        "title": _text(_pick(item, "title", "judul", "uraian", "description", "nama")),
        "risk": _text(_pick(item, "risk", "risiko", "tingkat_risiko", "risk_level")),
        "match": _gap_status(_pick(item, "match", "kesesuaian", "status")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
        "source": _text(_pick(item, "source", "sumber", "dokumen")),
    }


def _normalise_permit(item: dict[str, Any]) -> dict[str, Any]:
    kbli_raw = _pick(item, "kbli", "kode_kbli", "kbli_list", "business_classification")
    return {
        "name": _text(_pick(item, "name", "nama", "jenis_izin", "permit", "title")),
        "number": _text(_pick(item, "number", "nomor", "no", "nomor_izin", "nib")),
        "issued": _text(_pick(item, "issued", "tanggal_terbit", "tanggal", "date", "issued_at")),
        "expiry": _text(_pick(item, "expiry", "masa_berlaku", "berlaku_sampai", "valid_until", "expired_at")),
        "issuer": _text(_pick(item, "issuer", "instansi", "penerbit", "authority", "lembaga")),
        "status": _permit_status(_pick(item, "status", "keadaan")),
        "kbli": [_normalise_kbli(entry) for entry in _list(kbli_raw)],
        "match": _gap_status(_pick(item, "match", "kesesuaian", "compliance")),
        "gap": _text(_pick(item, "gap", "kekurangan", "selisih", "finding")),
        "severity": _severity(_pick(item, "severity", "risiko", "risk", "tingkat_risiko")),
        "action": _text(_pick(item, "action", "tindakan", "remediation", "next_step")),
        "source": _text(_pick(item, "source", "sumber", "dokumen", "document")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_covenant(item: dict[str, Any]) -> dict[str, Any]:
    restriction = _text(_pick(item, "restriction", "restriksi", "pembatasan", "covenant", "klausul", "summary"))
    return {
        "category": _text(_pick(item, "category", "kategori", "type", "jenis")),
        "document": _text(_pick(item, "document", "dokumen", "perjanjian", "agreement", "source")),
        "creditor": _text(_pick(item, "creditor", "kreditur", "bank", "lender")),
        "clause": _text(_pick(item, "clause", "pasal", "klausul", "article", "section")),
        "summary": _text(_pick(item, "summary", "ringkasan", "uraian", "description")),
        "restriction": restriction,
        "threshold": _text(_pick(item, "threshold", "ambang", "batas", "limit")),
        "actual": _text(_pick(item, "actual", "kondisi_aktual", "current", "posisi")),
        "status": _gap_status(_pick(item, "status", "kepatuhan", "compliance")),
        "severity": _severity(_pick(item, "severity", "risiko", "risk")),
        "action": _text(_pick(item, "action", "tindakan", "remediation")),
        "source": _text(_pick(item, "source", "sumber", "dokumen")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_loan(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "document": _text(_pick(item, "document", "dokumen", "perjanjian", "agreement", "name", "nama")),
        "creditor": _text(_pick(item, "creditor", "kreditur", "bank", "lender")),
        "facility": _text(_pick(item, "facility", "fasilitas", "jenis_fasilitas", "type")),
        "principal": _number(_pick(item, "principal", "plafon", "plafond", "nilai", "amount", "limit")),
        "balance": _number(_pick(item, "balance", "saldo", "outstanding", "baki_debet")),
        "currency": _text(_pick(item, "currency", "mata_uang"), "IDR"),
        "due": _text(_pick(item, "due", "jatuh_tempo", "maturity", "due_date")),
        "clause": _text(_pick(item, "clause", "pasal", "klausul")),
        "restriction": _text(_pick(item, "restriction", "restriksi", "pembatasan", "covenant")),
        "status": _gap_status(_pick(item, "status", "kepatuhan", "compliance")),
        "severity": _severity(_pick(item, "severity", "risiko", "risk")),
        "source": _text(_pick(item, "source", "sumber", "dokumen")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_litigation(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "case_number": _text(_pick(item, "case_number", "nomor_perkara", "register", "no_perkara", "number")),
        "court": _text(_pick(item, "court", "pengadilan", "forum", "institution")),
        "parties": _text(_pick(item, "parties", "para_pihak", "pihak", "party")),
        "role": _text(_pick(item, "role", "peran", "posisi", "capacity")),
        "subject": _text(_pick(item, "subject", "pokok_perkara", "objek", "matter", "description")),
        "value": _number(_pick(item, "value", "nilai", "nilai_gugatan", "amount", "claim")),
        "status": _litigation_status(_pick(item, "status", "keadaan", "stage")),
        "severity": _severity(_pick(item, "severity", "risiko", "risk")),
        "asset_impact": _text(_pick(item, "asset_impact", "dampak_aset", "asset", "aset")),
        "source": _text(_pick(item, "source", "sumber", "dokumen")),
        "action": _text(_pick(item, "action", "tindakan", "next_step")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_asset(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": _text(_pick(item, "name", "nama", "aset", "asset", "description")),
        "asset_type": _text(_pick(item, "asset_type", "jenis", "type", "kategori")),
        "ownership": _text(_pick(item, "ownership", "kepemilikan", "owner", "pemilik")),
        "certificate": _text(_pick(item, "certificate", "sertifikat", "bukti_kepemilikan", "certificate_number")),
        "encumbrance": _text(_pick(item, "encumbrance", "beban", "jaminan", "mortgage", "hak_tanggungan")),
        "dispute": _text(_pick(item, "dispute", "sengketa", "litigation", "status_sengketa")),
        "status": _gap_status(_pick(item, "status", "keadaan")),
        "severity": _severity(_pick(item, "severity", "risiko", "risk")),
        "source": _text(_pick(item, "source", "sumber", "dokumen")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_red_flag(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "severity": _severity(_pick(item, "severity", "risiko", "risk", "tingkat")),
        "issue": _text(_pick(item, "issue", "isu", "temuan", "finding", "title", "masalah")),
        "basis": _text(_pick(item, "basis", "dasar", "fakta", "document", "sumber", "evidence")),
        "consequence": _text(_pick(item, "consequence", "konsekuensi", "dampak", "impact", "risk_description")),
        "action": _text(_pick(item, "action", "tindakan", "recommendation", "mitigasi", "remediation")),
        "owner": _text(_pick(item, "owner", "pic", "penanggung_jawab")),
        "deadline": _text(_pick(item, "deadline", "tenggat", "due")),
    }


def _normalise_document(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": _text(_pick(item, "name", "nama", "filename", "file", "dokumen")),
        "category": _text(_pick(item, "category", "kategori", "type", "jenis"), "Lainnya"),
        "path": _text(_pick(item, "path", "lokasi", "file_path")),
        "status": _text(_pick(item, "status", "keadaan"), "Ditemukan"),
        "size": int(_number(_pick(item, "size", "bytes", "ukuran")) or 0) or None,
        "modified": _text(_pick(item, "modified", "diubah", "tanggal")),
        "notes": _text(_pick(item, "notes", "catatan", "keterangan")),
    }


def _normalise_corporate_action(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "date": _text(_pick(item, "date", "tanggal", "tahun")),
        "action": _text(_pick(item, "action", "peristiwa", "event", "description")),
        "document": _text(_pick(item, "document", "dokumen", "akta", "source")),
        "impact": _text(_pick(item, "impact", "dampak", "effect", "notes")),
    }


def normalise_analysis(raw: Any) -> Analysis:
    data = _unwrap_analysis(raw)
    company = _dict(_pick(data, "company", "entity", "entitas", "profil_perusahaan", "perusahaan"))
    deal = _dict(_pick(data, "deal", "transaction", "transaksi", "engagement", "mandat"))

    if not company and isinstance(data.get("target"), Mapping):
        company = dict(data["target"])

    company = {
        "name": _text(_pick(company, "name", "nama", "nama_perusahaan", "company_name")),
        "nib": _text(_pick(company, "nib", "nomor_nib", "nib_number")),
        "npwp": _text(_pick(company, "npwp", "tax_id", "nomor_npwp")),
        "domicile": _text(_pick(company, "domicile", "domisili", "alamat", "address")),
        "established": _text(_pick(company, "established", "tanggal_pendirian", "didirikan", "incorporation_date")),
        "deed_number": _text(_pick(company, "deed_number", "nomor_akta", "akta_pendirian", "deed")),
        "deed_date": _text(_pick(company, "deed_date", "tanggal_akta", "deed_date")),
        "notary": _text(_pick(company, "notary", "notaris", "notary_name")),
        "sk_number": _text(_pick(company, "sk_number", "nomor_sk", "sk_kemenkumham", "sk")),
        "sk_date": _text(_pick(company, "sk_date", "tanggal_sk")),
        "business": _text(_pick(company, "business", "kegiatan_usaha", "business_activity")),
        "notes": _text(_pick(company, "notes", "catatan", "keterangan")),
    }

    shareholders_raw = _pick(data, "shareholders", "pemegang_saham", "struktur_saham", "shareholding", "composition")
    if not shareholders_raw:
        shareholders_raw = _pick(company, "shareholders", "pemegang_saham", "struktur_saham")
    directors_raw = _pick(data, "directors", "direksi", "management", "pengurus")
    if not directors_raw:
        directors_raw = _pick(company, "directors", "direksi")
    commissioners_raw = _pick(data, "commissioners", "komisaris", "board_of_commissioners")
    if not commissioners_raw:
        commissioners_raw = _pick(company, "commissioners", "komisaris")

    permit_raw = _pick(data, "permits", "perizinan", "izin", "licenses", "oss_rba", "perizinan_oss")
    if not permit_raw:
        permit_raw = _pick(company, "permits", "perizinan", "izin")
    kbli_raw = _pick(data, "kbli", "kode_kbli", "kbli_list", "business_classification")
    if not kbli_raw:
        kbli_raw = _pick(company, "kbli", "kode_kbli")

    loans_raw = _pick(data, "loans", "utang", "utang_bank", "fasilitas_kredit", "perjanjian_kredit", "bank_debt")
    covenants_raw = _pick(data, "covenants", "negative_covenants", "restrictions", "klausul_negatif", "pembatasan")
    if not covenants_raw:
        covenants_raw = _pick(company, "covenants", "negative_covenants")

    permits = [_normalise_permit(item) for item in _dicts(permit_raw)]
    kbli_entries = [_normalise_kbli(item) for item in _list(kbli_raw)]
    for permit in permits:
        kbli_entries.extend(permit.get("kbli", []))

    # De-duplicate KBLI while keeping the first useful title/risk.
    seen_kbli: dict[str, dict[str, Any]] = {}
    for entry in kbli_entries:
        code = entry.get("code", "")
        key = code or f"{entry.get('title', '')}|{entry.get('source', '')}"
        if not key:
            continue
        if key not in seen_kbli:
            seen_kbli[key] = entry
        else:
            for field in ("title", "risk", "match", "notes", "source"):
                if not seen_kbli[key].get(field) and entry.get(field):
                    seen_kbli[key][field] = entry[field]

    documents = [_normalise_document(item) for item in _dicts(_pick(data, "documents", "dokumen", "data_room", "files"))]
    red_flags = [_normalise_red_flag(item) for item in _dicts(_pick(data, "red_flags", "redflags", "temuan_kritis", "critical_findings"))]

    findings_raw = _pick(data, "findings", "temuan", "issues", "checklist")
    for finding in _dicts(findings_raw):
        status = _gap_status(_pick(finding, "status", "kepatuhan", "compliance"))
        if status in {"Tidak Terpenuhi", "Tidak Ditemukan"}:
            red_flags.append(
                _normalise_red_flag(
                    {
                        "severity": _pick(finding, "severity", "risiko", "risk"),
                        "issue": _pick(finding, "issue", "finding", "temuan", "requirement", "persyaratan"),
                        "basis": _pick(finding, "basis", "source", "dokumen", "data"),
                        "consequence": _pick(finding, "consequence", "gap", "dampak"),
                        "action": _pick(finding, "action", "tindakan", "remediation"),
                    }
                )
            )

    corporate_actions = [
        _normalise_corporate_action(item)
        for item in _dicts(_pick(data, "corporate_actions", "aksi_korporasi", "history", "riwayat", "changes"))
    ]

    deal = {
        "name": _text(_pick(deal, "name", "nama", "deal_name", "transaction_name")),
        "client": _text(_pick(data, "client", "klien"), _text(_pick(deal, "client", "klien"))),
        "target": _text(_pick(data, "target_name", "target"), _text(_pick(deal, "target", "target_company")) or company["name"]),
        "reviewer": _text(_pick(deal, "reviewer", "peninjau", "prepared_by"), _text(_pick(data, "reviewer", "peninjau"))),
        "date": _text(_pick(deal, "date", "tanggal"), _text(_pick(data, "date", "tanggal"))),
        "objective": _text(_pick(deal, "objective", "tujuan", "scope", "lingkup")),
        "source": _text(_pick(deal, "source", "sumber")),
    }

    analysis = Analysis(
        deal=Deal.model_validate(deal),
        company=Company.model_validate(company),
        shareholders=[Shareholder.model_validate(_normalise_shareholder(item)) for item in _dicts(shareholders_raw)],
        directors=[Person.model_validate(_normalise_person(item, "Direktur")) for item in _dicts(directors_raw)],
        commissioners=[Person.model_validate(_normalise_person(item, "Komisaris")) for item in _dicts(commissioners_raw)],
        corporate_actions=corporate_actions,
        permits=[Permit.model_validate(item) for item in permits],
        kbli=[Kbli.model_validate(item) for item in seen_kbli.values()],
        loans=[Loan.model_validate(_normalise_loan(item)) for item in _dicts(loans_raw)],
        covenants=[Covenant.model_validate(_normalise_covenant(item)) for item in _dicts(covenants_raw)],
        litigation=[Litigation.model_validate(_normalise_litigation(item)) for item in _dicts(_pick(data, "litigation", "litigasi", "sengketa", "perkara", "disputes"))],
        assets=[Asset.model_validate(_normalise_asset(item)) for item in _dicts(_pick(data, "assets", "aset", "sertifikat_aset", "properties"))],
        red_flags=[RedFlag.model_validate(item) for item in red_flags],
        documents=[DocumentInfo.model_validate(item) for item in documents],
    )
    return analysis


# ---------------------------------------------------------------------------
# Data room directory inventory
# ---------------------------------------------------------------------------


CATEGORY_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("Litigasi / Sengketa", ("gugatan", "somasi", "putusan", "perkara", "litigasi", "sengketa", "arbitrase", "banding", "kasasi")),
    ("Perjanjian Kredit / Utang Bank", ("kredit", "fasilitas", "pinjaman", "loan", "utang", "debenture", "bank")),
    ("SK Kemenkumham / AHU", ("kemenkumham", "ahu", "sk menkum", "penerimaan pemberitahuan", "pengesahan", "sk menteri")),
    ("Akta PT", ("akta", "pendirian", "anggaran dasar", "deed", "articles of association", "ad perusahaan")),
    ("NIB / OSS-RBA", ("nib", "oss", "perizinan berusaha", "kbli", "izin usaha", "sertifikat standar", "izin sektor")),
    ("Aset / Sertifikat", ("sertifikat", "shm", "hgb", "tanah", "bangunan", "aset", "asset", "hak milik", "hak guna")),
    ("Kontrak Material", ("perjanjian", "kontrak", "agreement", "mou", "pks", "side letter", "addendum")),
    ("Korporasi Lain", ("rups", "saham", "direksi", "komisaris", "laporan", "npwp", "ktp", "passport")),
]

COVENANT_KEYWORDS: list[tuple[str, tuple[str, ...], str]] = [
    ("Perubahan Kepemilikan", ("perubahan kepemilikan", "change of control", "pengalihan saham", "transfer saham"), "HIGH"),
    ("Pembagian Dividen", ("pembagian dividen", "dividen", "dividend", "distribution"), "MEDIUM"),
    ("Penjaminan Aset", ("penjaminan aset", "jaminan", "agunan", "pledge", "security interest", "hak tanggungan", "fidusia"), "HIGH"),
    ("Utang Tambahan", ("utang tambahan", "additional indebtedness", "pinjaman baru", "additional debt"), "MEDIUM"),
    ("Cross Default", ("cross default", "wanprestasi silang", "cross-default"), "HIGH"),
    ("Pembatasan Investasi", ("pembatasan investasi", "capital expenditure", "capex", "investasi"), "MEDIUM"),
    ("Pembatasan Pengalihan Aset", ("pengalihan aset", "transfer aset", "sale of assets", "divestasi"), "HIGH"),
]

DATE_PATTERN = re.compile(
    r"\b(\d{1,2}[-/]\d{1,2}[-/]\d{4}|\d{1,2}\s+(?:Januari|Februari|Maret|April|Mei|Juni|Juli|Agustus|September|Oktober|November|Desember)\s+\d{4})\b",
    re.IGNORECASE,
)
NIB_PATTERN = re.compile(r"\bNIB\b[^\d]{0,40}(\d{13})", re.IGNORECASE)
NIB_ANY_PATTERN = re.compile(r"\b(\d{13})\b")
KBLI_PATTERN = re.compile(r"\bKBLI\b[^\d]{0,30}(\d{5})", re.IGNORECASE)
KBLI_ANY_PATTERN = re.compile(r"\b(\d{5})\b")
DEED_PATTERN = re.compile(
    r"(?:Akta|Deed)\s*(?:Pendirian|Perubahan|Pernyataan Keputusan Rapat)?\s*(?:Nomor|No\.?)\s*[:.]?\s*([0-9]{1,7})",
    re.IGNORECASE,
)
SK_PATTERN = re.compile(
    r"(?:Surat Keputusan|SK|Keputusan)\s*(?:Menteri|Menkumham)?\s*(?:Nomor|No\.?)\s*[:.]?\s*([A-Z0-9./\-]{3,60})",
    re.IGNORECASE,
)
COMPANY_PATTERN = re.compile(
    r"\bPT[.\s]+([A-Z][A-Za-z0-9&.'\-]*(?:\s+[A-Z][A-Za-z0-9&.'\-]*){0,7})"
)
BANK_PATTERN = re.compile(r"\bBank\s+([A-Z][A-Za-z0-9&.\-]+(?:\s+[A-Z][A-Za-z0-9&.\-]+){0,4})")
CERTIFICATE_PATTERN = re.compile(
    r"(?:Sertifikat|SHM|HGB|Hak Milik|Hak Guna Bangunan)\s*(?:Nomor|No\.?)?\s*[:.]?\s*([A-Z0-9./\-]{2,50})",
    re.IGNORECASE,
)
CASE_PATTERN = re.compile(
    r"(?:Nomor Perkara|No\.?\s*Perkara)\s*[:.]?\s*([0-9]{1,4}/[A-Za-z0-9./\-]+)",
    re.IGNORECASE,
)
COURT_PATTERN = re.compile(r"\b(Pengadilan Negeri\s+[A-Z][A-Za-z .\-]{2,50})")
AMOUNT_PATTERN = re.compile(r"\bRp\.?\s*[0-9][0-9.,]*", re.IGNORECASE)


def _classify_file(path: Path) -> str:
    name = path.name.lower().replace("_", " ").replace("-", " ")
    parent = path.parent.name.lower().replace("_", " ").replace("-", " ")
    haystack = f"{parent} {name}"
    for category, keywords in CATEGORY_RULES:
        if any(keyword in haystack for keyword in keywords):
            return category
    return "Lainnya"


def _decode_text_file(path: Path) -> str:
    raw = path.read_bytes()[: MAX_TEXT_PER_FILE * 4]
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw.decode(encoding)[:MAX_TEXT_PER_FILE]
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")[:MAX_TEXT_PER_FILE]


def _docx_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            payload = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile, OSError):
        return ""
    try:
        root = ElementTree.fromstring(payload)
    except ElementTree.ParseError:
        return ""
    texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
    return "\n".join(texts)[:MAX_TEXT_PER_FILE]


def _xlsx_text(path: Path) -> str:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ""
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(sheet.title)
            for row in sheet.iter_rows(max_row=500, max_col=30, values_only=True):
                values = [_text(value) for value in row if _text(value)]
                if values:
                    lines.append(" | ".join(values))
                if sum(len(line) for line in lines) > MAX_TEXT_PER_FILE:
                    return "\n".join(lines)[:MAX_TEXT_PER_FILE]
    finally:
        workbook.close()
    return "\n".join(lines)[:MAX_TEXT_PER_FILE]


def _pdf_text(path: Path) -> tuple[str, str]:
    try:
        import pymupdf as fitz  # type: ignore
    except ImportError:
        try:
            import fitz  # type: ignore
        except ImportError:
            return "", "PyMuPDF tidak tersedia; PDF tidak diekstrak."
    try:
        document = fitz.open(path)
    except Exception as exc:
        return "", f"PDF tidak dapat dibuka: {exc}"
    if getattr(document, "needs_pass", False):
        document.close()
        return "", "PDF terenkripsi dan memerlukan kata sandi."
    parts: list[str] = []
    try:
        for index in range(min(document.page_count, 200)):
            parts.append(document.load_page(index).get_text("text"))
            if sum(len(part) for part in parts) > MAX_TEXT_PER_FILE:
                break
    finally:
        document.close()
    text = "\n".join(parts)[:MAX_TEXT_PER_FILE]
    if len(text.strip()) < 80:
        return text, "PDF tampak hasil pemindaian; jalankan OCR untuk membaca isinya."
    return text, ""


def _read_document_text(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm"}:
        return _decode_text_file(path), ""
    if suffix == ".docx":
        return _docx_text(path), ""
    if suffix == ".xlsx":
        return _xlsx_text(path), ""
    if suffix == ".pdf":
        return _pdf_text(path)
    if suffix == ".doc":
        return "", "Format .doc lama tidak dibaca langsung; konversi ke .docx."
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        return "", "Berkas gambar memerlukan OCR."
    return "", "Jenis berkas tidak diekstrak."


def _extract_facts(text: str) -> dict[str, Any]:
    facts: dict[str, Any] = {
        "nib": "",
        "kbli": [],
        "company": "",
        "deed_number": "",
        "sk_number": "",
        "dates": [],
        "banks": [],
        "certificates": [],
        "cases": [],
        "courts": [],
        "amounts": [],
        "covenants": [],
    }
    if not text:
        return facts
    nib = NIB_PATTERN.search(text)
    if nib:
        facts["nib"] = nib.group(1)
    else:
        fallback = NIB_ANY_PATTERN.search(text)
        if fallback:
            facts["nib"] = fallback.group(1)
    for match in KBLI_PATTERN.finditer(text):
        facts["kbli"].append(match.group(1))
    if not facts["kbli"]:
        for match in KBLI_ANY_PATTERN.finditer(text):
            context = text[max(0, match.start() - 60): match.end() + 60].lower()
            if "kbli" in context or "klasifikasi baku" in context:
                facts["kbli"].append(match.group(1))
    company = COMPANY_PATTERN.search(text)
    if company:
        facts["company"] = f"PT {company.group(1).strip()}"
    deed = DEED_PATTERN.search(text)
    if deed:
        facts["deed_number"] = deed.group(1)
    sk = SK_PATTERN.search(text)
    if sk:
        facts["sk_number"] = sk.group(1)
    facts["dates"] = list(dict.fromkeys(match.group(1) for match in DATE_PATTERN.finditer(text)))[:10]
    facts["banks"] = list(dict.fromkeys(match.group(1).strip() for match in BANK_PATTERN.finditer(text)))[:5]
    facts["certificates"] = list(dict.fromkeys(match.group(1).strip() for match in CERTIFICATE_PATTERN.finditer(text)))[:10]
    facts["cases"] = list(dict.fromkeys(match.group(1).strip() for match in CASE_PATTERN.finditer(text)))[:10]
    facts["courts"] = list(dict.fromkeys(match.group(1).strip() for match in COURT_PATTERN.finditer(text)))[:5]
    facts["amounts"] = list(dict.fromkeys(match.group(0).strip() for match in AMOUNT_PATTERN.finditer(text)))[:5]
    lowered = text.lower()
    for category, keywords, _severity_value in COVENANT_KEYWORDS:
        for keyword in keywords:
            position = lowered.find(keyword)
            if position < 0:
                continue
            excerpt = text[max(0, position - 180): position + 320].replace("\n", " ").strip()
            facts["covenants"].append(
                {
                    "category": category,
                    "restriction": keyword,
                    "summary": excerpt[:500],
                }
            )
            break
    return facts


def _snippet(text: str, start: int, end: int) -> str:
    return text[max(0, start - 160): min(len(text), end + 260)].replace("\n", " ").strip()


def analyse_directory(
    directory: Path,
    *,
    client: str = "",
    target: str = "",
    reviewer: str = "",
) -> Analysis:
    if not directory.is_dir():
        raise DueDiligenceError(f"{directory} is not a directory.")
    files: list[Path] = []
    for root, dirs, names in os.walk(directory, followlinks=False):
        dirs[:] = sorted(
            name for name in dirs if not name.startswith(".") and name not in {"node_modules", "__pycache__"}
        )
        for name in sorted(names):
            if name.startswith("."):
                continue
            path = Path(root) / name
            if path.is_symlink() or not path.is_file():
                continue
            files.append(path)
            if len(files) >= MAX_DATA_ROOM_FILES:
                break
        if len(files) >= MAX_DATA_ROOM_FILES:
            break
    if not files:
        raise DueDiligenceError(f"{directory} contains no readable files.")

    documents: list[DocumentInfo] = []
    warnings: list[str] = []
    company = {"name": target, "nib": "", "deed_number": "", "sk_number": "", "notes": ""}
    kbli: list[Kbli] = []
    permits: list[Permit] = []
    loans: list[Loan] = []
    covenants: list[Covenant] = []
    litigation: list[Litigation] = []
    assets: list[Asset] = []
    red_flags: list[RedFlag] = []
    seen_categories: set[str] = set()
    seen_kbli: set[str] = set()
    seen_covenants: set[tuple[str, str]] = set()
    seen_cases: set[str] = set()

    for path in files:
        category = _classify_file(path)
        seen_categories.add(category)
        relative = os.path.relpath(path, directory)
        stat = path.stat()
        text, warning = _read_document_text(path)
        facts = _extract_facts(text)
        if warning:
            warnings.append(f"{relative}: {warning}")
        notes = warning
        if facts.get("nib"):
            company["nib"] = company["nib"] or facts["nib"]
        if facts.get("company") and not company["name"]:
            company["name"] = facts["company"]
        if facts.get("deed_number"):
            company["deed_number"] = company["deed_number"] or facts["deed_number"]
        if facts.get("sk_number"):
            company["sk_number"] = company["sk_number"] or facts["sk_number"]

        documents.append(
            DocumentInfo(
                name=path.name,
                category=category,
                path=relative,
                status="Ditemukan" if text or not warning else "Perlu Klarifikasi",
                size=stat.st_size,
                modified=datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).strftime("%Y-%m-%d"),
                notes=notes,
            )
        )

        for code in facts.get("kbli", []):
            if code in seen_kbli:
                continue
            seen_kbli.add(code)
            kbli.append(
                Kbli(
                    code=code,
                    title="",
                    risk="",
                    match="Perlu Klarifikasi",
                    notes="Kode ditemukan di dokumen; cocokkan dengan kegiatan usaha riil.",
                    source=relative,
                )
            )

        if category == "NIB / OSS-RBA" and not any(permit.name.upper().startswith("NIB") for permit in permits):
            permits.append(
                Permit(
                    name="NIB / Perizinan Berusaha",
                    number=facts.get("nib", ""),
                    issuer="OSS-RBA",
                    status="Perlu Klarifikasi",
                    kbli=[Kbli.model_validate(_normalise_kbli(code)) for code in facts.get("kbli", [])],
                    match="Perlu Klarifikasi",
                    gap="Status aktif, izin turunan, dan pemenuhan kewajiban belum diverifikasi.",
                    severity="HIGH",
                    action="Verifikasi NIB dan izin sektor pada portal OSS-RBA; unduh riwayat perizinan.",
                    source=relative,
                    notes="Keberadaan dokumen tidak membuktikan status izin aktif.",
                )
            )

        if category == "Perjanjian Kredit / Utang Bank":
            bank = facts.get("banks", [""])[0] if facts.get("banks") else ""
            amount = facts.get("amounts", [""])[0] if facts.get("amounts") else ""
            loans.append(
                Loan(
                    document=path.stem,
                    creditor=bank,
                    facility="",
                    principal=_number(amount),
                    due="",
                    status="Perlu Klarifikasi",
                    severity="HIGH",
                    source=relative,
                    notes="Fasilitas, saldo, jadwal, dan covenant belum diverifikasi.",
                )
            )
            for covenant in facts.get("covenants", []):
                key = (covenant["category"], relative)
                if key in seen_covenants:
                    continue
                seen_covenants.add(key)
                covenants.append(
                    Covenant(
                        category=covenant["category"],
                        document=path.stem,
                        creditor=bank,
                        summary=covenant.get("summary", ""),
                        restriction=covenant.get("restriction", ""),
                        status="Perlu Klarifikasi",
                        severity="HIGH" if covenant["category"] in {"Perubahan Kepemilikan", "Penjaminan Aset", "Cross Default"} else "MEDIUM",
                        action="Baca klausul lengkap, catat ambang, dan uji kepatuhan terhadap kondisi aktual.",
                        source=relative,
                    )
                )

        if category == "Litigasi / Sengketa":
            cases = facts.get("cases", [])
            court = facts.get("courts", [""])[0] if facts.get("courts") else ""
            case_number = cases[0] if cases else ""
            key = case_number or relative
            if key not in seen_cases:
                seen_cases.add(key)
                litigation.append(
                    Litigation(
                        case_number=case_number,
                        court=court,
                        subject=path.stem,
                        status="Perlu Klarifikasi",
                        severity="HIGH",
                        source=relative,
                        action="Verifikasi register perkara, para pihak, tahap, dan nilai gugatan.",
                        notes="Dokumen ditemukan; status perkara belum diverifikasi.",
                    )
                )

        if category == "Aset / Sertifikat":
            certificates = facts.get("certificates", [])
            assets.append(
                Asset(
                    name=path.stem,
                    certificate=certificates[0] if certificates else "",
                    ownership=company["name"],
                    encumbrance="Perlu Klarifikasi",
                    dispute="Tidak Ditemukan",
                    status="Perlu Klarifikasi",
                    severity="MEDIUM",
                    source=relative,
                    notes="Cek keaslian sertifikat, beban hak tanggungan, dan kesesuaian tata ruang.",
                )
            )

    missing = {
        "Akta PT": ("Berkas akta pendirian/perubahan tidak ditemukan di data room.", "HIGH"),
        "SK Kemenkumham / AHU": ("SK Kemenkumham atau surat penerimaan pemberitahuan AHU tidak ditemukan.", "HIGH"),
        "NIB / OSS-RBA": ("Dokumen NIB/OSS-RBA tidak ditemukan di data room.", "HIGH"),
        "Perjanjian Kredit / Utang Bank": ("Perjanjian kredit/utang bank tidak ditemukan; kewajiban dan covenant belum dapat dinilai.", "MEDIUM"),
        "Aset / Sertifikat": ("Dokumen kepemilikan aset tidak ditemukan; status jaminan belum dapat dinilai.", "MEDIUM"),
        "Litigasi / Sengketa": ("Tidak ada berkas litigasi; konfirmasi apakah benar tidak ada perkara berjalan.", "LOW"),
    }
    for category, (message, severity_value) in missing.items():
        if category not in seen_categories:
            red_flags.append(
                RedFlag(
                    severity=severity_value,
                    issue=message,
                    basis=f"Inventaris data room: {len(files)} berkas diperiksa.",
                    consequence="Ketidakpastian hukum dan risiko penutupan transaksi.",
                    action=f"Minta {category} terbaru dan verifikasi ke sumber resmi.",
                )
            )

    if company["nib"]:
        red_flags.append(
            RedFlag(
                severity="MEDIUM",
                issue="Nomor NIB ditemukan, tetapi status aktif dan izin turunannya belum diverifikasi.",
                basis=f"NIB {company['nib']} disebut dalam dokumen data room.",
                consequence="Risiko beroperasi tanpa izin yang sah bila NIB atau izin sektor tidak aktif.",
                action="Verifikasi NIB, KBLI, dan izin sektor pada OSS-RBA; simpan tangkapan layar.",
            )
        )
    if litigation:
        red_flags.append(
            RedFlag(
                severity="HIGH",
                issue=f"{len(litigation)} perkara/sengketa teridentifikasi dari berkas data room.",
                basis=", ".join(item.source for item in litigation[:5]),
                consequence="Potensi kewajiban, sitaan aset, atau hambatan closing.",
                action="Minta status register perkara dan opini litigasi; nilai provisi.",
            )
        )
    if covenants:
        red_flags.append(
            RedFlag(
                severity="HIGH",
                issue=f"{len(covenants)} indikasi negative covenant ditemukan pada dokumen kredit.",
                basis=", ".join(sorted({item.category for item in covenants})),
                consequence="Pelanggaran covenant dapat memicu percepatan jatuh tempo atau cross default.",
                action="Uji kepatuhan tiap covenant dan minta waiver bila diperlukan.",
            )
        )

    company_model = Company.model_validate(
        {
            "name": company["name"] or target or directory.name,
            "nib": company["nib"],
            "deed_number": company["deed_number"],
            "sk_number": company["sk_number"],
            "notes": "Diekstrak otomatis dari nama dan isi dokumen; wajib diverifikasi terhadap dokumen asli.",
        }
    )
    analysis = Analysis(
        deal=Deal(
            name=f"Due diligence — {directory.name}",
            client=client,
            target=target or company_model.name,
            reviewer=reviewer,
            date=date.today().isoformat(),
            source=str(directory),
        ),
        company=company_model,
        permits=permits,
        kbli=kbli,
        loans=loans,
        covenants=covenants,
        litigation=litigation,
        assets=assets,
        red_flags=red_flags,
        documents=documents,
    )
    analysis.__dict__["_warnings"] = warnings
    return analysis


# ---------------------------------------------------------------------------
# Workbook styling
# ---------------------------------------------------------------------------


def _title(ws: Worksheet, title: str, subtitle: str, last_col: int, meta: str) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    cell = ws.cell(row=2, column=1, value=subtitle)
    cell.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    cell = ws.cell(row=3, column=1, value=meta)
    cell.font = Font(name="Calibri", size=9, color="595959")
    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 16


def _section(ws: Worksheet, row: int, label: str, last_col: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cell.border = BORDER
    ws.row_dimensions[row].height = 20
    return row + 1


def _headers(ws: Worksheet, row: int, headers: Sequence[str]) -> int:
    for column, value in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30
    return row + 1


def _value_cell(cell: Any, value: Any, *, number_format: str | None = None) -> None:
    cell.value = "" if value is None else value
    cell.font = Font(name="Calibri", size=9)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if number_format:
        cell.number_format = number_format


def _paint_severity(cell: Any, value: Any) -> None:
    text = _text(value).upper()
    if text == "HIGH":
        cell.fill = PatternFill("solid", fgColor=RED_FILL)
        cell.font = Font(name="Calibri", size=9, bold=True, color=DARK_RED)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    elif text == "MEDIUM":
        cell.fill = PatternFill("solid", fgColor=AMBER_FILL)
        cell.font = Font(name="Calibri", size=9, bold=True, color=DARK_AMBER)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    elif text == "LOW":
        cell.fill = PatternFill("solid", fgColor=GREEN_FILL)
        cell.font = Font(name="Calibri", size=9, bold=True, color=DARK_GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _paint_status(cell: Any, value: Any) -> None:
    text = _text(value)
    if text in {"Tidak Terpenuhi", "Tidak Ditemukan", "Kedaluwarsa", "Dicabut"}:
        cell.fill = PatternFill("solid", fgColor=RED_FILL)
        cell.font = Font(name="Calibri", size=9, bold=True, color=DARK_RED)
    elif text in {"Perlu Klarifikasi", "Akan Berakhir"}:
        cell.fill = PatternFill("solid", fgColor=AMBER_FILL)
        cell.font = Font(name="Calibri", size=9, bold=True, color=DARK_AMBER)
    elif text in {"Terpenuhi", "Berlaku", "Selesai"}:
        cell.fill = PatternFill("solid", fgColor=GREEN_FILL)
        cell.font = Font(name="Calibri", size=9, bold=True, color=DARK_GREEN)


def _rows(
    ws: Worksheet,
    row: int,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    widths: Sequence[float],
    severity_col: int | None = None,
    status_col: int | None = None,
    money_cols: Sequence[int] = (),
    percent_cols: Sequence[int] = (),
) -> tuple[int, int]:
    header_row = row
    row = _headers(ws, row, headers)
    last_data_row = row - 1
    data = list(rows)[:MAX_WORKBOOK_ROWS]
    for index, values in enumerate(data):
        last_data_row = row
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column)
            number_format = None
            if column in money_cols:
                number_format = "#,##0"
            elif column in percent_cols:
                number_format = "0.00%"
                if isinstance(value, (int, float)):
                    value = float(value) / 100.0
            _value_cell(cell, value, number_format=number_format)
            if column == severity_col:
                _paint_severity(cell, value)
            elif column == status_col:
                _paint_status(cell, value)
            elif index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=GREY)
        ws.row_dimensions[row].height = 32 if any(len(_text(v)) > 70 for v in values) else 22
        row += 1
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    if last_data_row >= header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_data_row}"
        ws.freeze_panes = f"A{header_row + 1}"
    return row, last_data_row


def _note(ws: Worksheet, row: int, message: str, last_col: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=message)
    cell.font = Font(name="Calibri", size=8, italic=True, color="595959")
    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    cell.border = BORDER
    ws.row_dimensions[row].height = 30
    return row + 1


def _finish_sheet(ws: Worksheet, last_row: int, last_col: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "RAHASIA — Dokumen kerja internal. Dihasilkan oleh Legal Agent."
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = "Halaman &P dari &N"
    ws.oddFooter.right.size = 8
    ws.print_title_rows = "1:4"
    if last_row < 4:
        ws.row_dimensions[4].height = 10


def _kv_rows(ws: Worksheet, row: int, values: Sequence[tuple[str, Any]]) -> int:
    for label, value in values:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        left = ws.cell(row=row, column=1, value=label)
        left.font = Font(name="Calibri", size=9, bold=True, color=NAVY)
        left.fill = PatternFill("solid", fgColor=PALE_BLUE)
        left.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        left.border = BORDER
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        right = ws.cell(row=row, column=3, value="" if value is None else value)
        right.font = Font(name="Calibri", size=9)
        right.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        right.border = BORDER
        ws.row_dimensions[row].height = 20
        row += 1
    return row


# ---------------------------------------------------------------------------
# Workbook assembly
# ---------------------------------------------------------------------------


def _percent(value: float | None) -> str:
    return f"{value:.2f}%" if value is not None else ""


def _money(value: float | None) -> float | None:
    return value


def _overall_risk(analysis: Analysis) -> str:
    severities = [flag.severity for flag in analysis.red_flags]
    severities.extend(covenant.severity for covenant in analysis.covenants)
    severities.extend(item.severity for item in analysis.litigation)
    severities.extend(permit.severity for permit in analysis.permits if permit.status in {"Kedaluwarsa", "Dicabut", "Tidak Ditemukan"})
    if "HIGH" in severities:
        return "HIGH"
    if "MEDIUM" in severities:
        return "MEDIUM"
    if "LOW" in severities:
        return "LOW"
    return "PERLU KLARIFIKASI"


def _count_severity(analysis: Analysis) -> dict[str, int]:
    counts = {"HIGH": 0, "MEDIUM": 0, "LOW": 0}
    for flag in analysis.red_flags:
        counts[flag.severity] = counts.get(flag.severity, 0) + 1
    return counts


def build_workbook(analysis: Analysis, output: Path) -> dict[str, Any]:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_SUMMARY

    meta = (
        f"Target: {analysis.company.name or analysis.deal.target or 'tidak ditentukan'}  |  "
        f"Klien: {analysis.deal.client or 'tidak ditentukan'}  |  "
        f"Tanggal review: {analysis.deal.date or date.today().isoformat()}  |  "
        f"Reviewer: {analysis.deal.reviewer or 'tidak ditentukan'}"
    )

    # Sheet 1 -----------------------------------------------------------------
    _title(ws, "EXECUTIVE SUMMARY & RED FLAGS", "Matriks Due Diligence Korporasi, Perizinan, Pembiayaan, dan Litigasi", 7, meta)
    severity_counts = _count_severity(analysis)
    row = _section(ws, 5, "A. INFORMASI TRANSAKSI & LINGKUP", 7)
    row = _kv_rows(
        ws,
        row,
        [
            ("Nama Transaksi / Proyek", analysis.deal.name),
            ("Klien", analysis.deal.client),
            ("Perusahaan Target", analysis.company.name or analysis.deal.target),
            ("NIB", analysis.company.nib),
            ("Kedudukan / Domisili", analysis.company.domicile),
            ("Tujuan Review", analysis.deal.objective),
            ("Sumber Data", analysis.deal.source or "Data room / analisis JSON"),
            ("Reviewer", analysis.deal.reviewer),
        ],
    )
    row += 1
    row = _section(ws, row, "B. RINGKASAN EKSEKUTIF", 7)
    row = _kv_rows(
        ws,
        row,
        [
            ("Jumlah Dokumen Diperiksa", len(analysis.documents)),
            ("Pemegang Saham Tercatat", len(analysis.shareholders)),
            ("Direksi / Komisaris Tercatat", len(analysis.directors) + len(analysis.commissioners)),
            ("Perizinan Tercatat", len(analysis.permits)),
            ("Kode KBLI Teridentifikasi", len(analysis.kbli)),
            ("Fasilitas Kredit / Utang Bank", len(analysis.loans)),
            ("Negative Covenant Teridentifikasi", len(analysis.covenants)),
            ("Perkara Litigasi / Sengketa", len(analysis.litigation)),
            ("Aset Tercatat", len(analysis.assets)),
            ("Red Flags (HIGH / MEDIUM / LOW)", f"{severity_counts.get('HIGH', 0)} / {severity_counts.get('MEDIUM', 0)} / {severity_counts.get('LOW', 0)}"),
            ("Tingkat Risiko Keseluruhan", _overall_risk(analysis)),
        ],
    )
    row += 1
    row = _section(ws, row, "C. RED FLAGS & TEMUAN KRITIS", 7)
    flag_rows = [
        [
            index,
            flag.severity,
            flag.issue,
            flag.basis,
            flag.consequence,
            flag.action,
            flag.owner,
            flag.deadline,
        ]
        for index, flag in enumerate(analysis.red_flags, start=1)
    ] or [[1, "MEDIUM", "Belum ada red flag yang dilaporkan.", "Tidak ada temuan eksplisit dalam analisis.", "Risiko belum terukur.", "Lengkapi analisis dan verifikasi dokumen.", "", ""]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Severity", "Isu / Temuan", "Dasar Fakta / Dokumen", "Konsekuensi Hukum", "Tindakan yang Disarankan", "PIC", "Tenggat"],
        flag_rows,
        widths=[5, 10, 42, 40, 35, 40, 18, 14],
        severity_col=2,
    )
    row += 1
    row = _section(ws, row, "D. TEMUAN KUNCI & CATATAN METODOLOGI", 7)
    key_findings = [
        f"• {flag.issue}" + (f" ({flag.severity})" if flag.severity else "")
        for flag in analysis.red_flags[:8]
    ]
    if not key_findings:
        key_findings = ["• Tidak ada temuan kunci yang dilaporkan; dokumen dan analisis masih perlu dilengkapi."]
    key_findings.append("")
    key_findings.append(
        "• Matriks ini adalah alat kerja, bukan opini hukum final. Setiap temuan wajib diverifikasi terhadap dokumen asli dan sumber resmi (AHU, OSS-RBA, pengadilan)."
    )
    key_findings.append(
        "• Status 'Tidak Ditemukan' berarti dokumen tidak ada di data room, bukan berarti persyaratan tidak berlaku."
    )
    for line in key_findings:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        cell = ws.cell(row=row, column=1, value=line)
        cell.font = Font(name="Calibri", size=9)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        cell.border = BORDER
        ws.row_dimensions[row].height = 28 if len(line) > 110 else 18
        row += 1
    _finish_sheet(ws, row, 7)

    # Sheet 2 -----------------------------------------------------------------
    ws = wb.create_sheet(SHEET_OWNERSHIP)
    _title(ws, "STRUKTUR SAHAM & DIREKSI HISTORIS", "Kepemilikan, pengurus, dan perubahan anggaran dasar", 9, meta)
    row = 5
    row = _section(ws, row, "A. STRUKTUR KEPEMILIKAN SAHAM", 9)
    shareholder_rows = [
        [
            index,
            item.name,
            item.share_class,
            item.shares,
            item.nominal,
            item.percentage,
            item.since,
            item.source,
            item.notes,
        ]
        for index, item in enumerate(analysis.shareholders, start=1)
    ] or [[1, "Tidak ditemukan di data room", "", "", "", "", "", "", "Perlu klarifikasi; minta daftar pemegang saham terbaru."]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Pemegang Saham", "Jenis Saham", "Jumlah Saham", "Nominal per Saham (Rp)", "Persentase", "Sejak", "Sumber Dokumen", "Catatan"],
        shareholder_rows,
        widths=[5, 30, 15, 13, 18, 11, 14, 30, 30],
        money_cols=(5,),
        percent_cols=(6,),
    )
    row += 1
    row = _section(ws, row, "B. DIREKSI & KOMISARIS (HISTORIS)", 9)
    people = list(analysis.directors) + list(analysis.commissioners)
    person_rows = [
        [index, item.name, item.position, item.period, item.basis, item.status, item.source]
        for index, item in enumerate(people, start=1)
    ] or [[1, "Tidak ditemukan di data room", "", "", "", "Perlu Klarifikasi", ""]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Nama", "Jabatan", "Periode", "Dasar Pengangkatan", "Status", "Sumber Dokumen"],
        person_rows,
        widths=[5, 30, 22, 18, 35, 18, 30],
        status_col=6,
    )
    row += 1
    row = _section(ws, row, "C. AKSI KORPORASI & PERUBAHAN ANGGARAN DASAR", 9)
    action_rows = [
        [index, item.get("date", ""), item.get("action", ""), item.get("document", ""), item.get("impact", "")]
        for index, item in enumerate(analysis.corporate_actions, start=1)
    ] or [[1, "", "Tidak ada riwayat aksi korporasi yang dilaporkan.", "", "Lengkapi dengan seluruh akta perubahan dan SK AHU."]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Tanggal", "Peristiwa / Aksi Korporasi", "Dokumen / Akta", "Dampak terhadap Struktur"],
        action_rows,
        widths=[5, 18, 50, 35, 45],
    )
    row = _note(
        ws,
        row,
        "Catatan: struktur historis hanya dapat dipastikan dengan seluruh akta perubahan, SK Kemenkumham/AHU, dan daftar pemegang saham yang konsisten. Jangan menyimpulkan kepemilikan akhir dari satu dokumen saja.",
        9,
    )
    _finish_sheet(ws, row, 9)

    # Sheet 3 -----------------------------------------------------------------
    ws = wb.create_sheet(SHEET_PERMITS)
    _title(ws, "PERIZINAN & OSS-RBA", "Checklist perizinan berusaha, KBLI, status, dan instansi", 13, meta)
    row = 5
    status_counts: dict[str, int] = {}
    for permit in analysis.permits:
        status_counts[permit.status] = status_counts.get(permit.status, 0) + 1
    row = _section(ws, row, "A. RINGKASAN STATUS PERIZINAN", 13)
    row = _kv_rows(
        ws,
        row,
        [
            ("Total Perizinan Tercatat", len(analysis.permits)),
            ("Berlaku", status_counts.get("Berlaku", 0)),
            ("Akan Berakhir", status_counts.get("Akan Berakhir", 0)),
            ("Kedaluwarsa / Dicabut", status_counts.get("Kedaluwarsa", 0) + status_counts.get("Dicabut", 0)),
            ("Perlu Klarifikasi", status_counts.get("Perlu Klarifikasi", 0)),
            ("Kode KBLI Teridentifikasi", len(analysis.kbli)),
        ],
    )
    row += 1
    row = _section(ws, row, "B. CHECKLIST PERIZINAN & OSS-RBA", 13)
    permit_rows = []
    for index, permit in enumerate(analysis.permits, start=1):
        permit_rows.append(
            [
                index,
                permit.name,
                permit.number,
                permit.issued,
                permit.expiry,
                permit.issuer,
                permit.status,
                ", ".join(kbli.code for kbli in permit.kbli if kbli.code),
                ", ".join(sorted({kbli.risk for kbli in permit.kbli if kbli.risk})),
                permit.match,
                permit.gap,
                permit.severity,
                permit.action,
                permit.source,
            ]
        )
    if not permit_rows:
        permit_rows = [[1, "Tidak ditemukan data perizinan", "", "", "", "", "Perlu Klarifikasi", "", "", "Perlu Klarifikasi", "Minta NIB dan seluruh izin sektor.", "HIGH", "Verifikasi ke OSS-RBA.", ""]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Jenis Izin", "Nomor", "Tanggal Terbit", "Masa Berlaku", "Instansi / Penerbit", "Status", "KBLI", "Risiko KBLI", "Kesesuaian", "Gap / Temuan", "Severity", "Tindakan", "Sumber"],
        permit_rows,
        widths=[5, 26, 20, 13, 13, 20, 15, 14, 12, 15, 35, 10, 35, 26],
        severity_col=12,
        status_col=7,
    )
    row += 1
    row = _section(ws, row, "C. KBLI & KESESUAIAN KEGIATAN USAHA", 13)
    kbli_rows = [
        [index, item.code, item.title, item.risk, item.match, item.notes, item.source]
        for index, item in enumerate(analysis.kbli, start=1)
    ] or [[1, "", "Tidak ada kode KBLI yang teridentifikasi.", "", "Perlu Klarifikasi", "Minta NIB dan cocokkan KBLI dengan kegiatan operasional.", ""]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Kode KBLI", "Judul KBLI", "Tingkat Risiko", "Kesesuaian dengan Kegiatan", "Catatan", "Sumber"],
        kbli_rows,
        widths=[5, 14, 40, 16, 24, 40, 28],
        status_col=5,
    )
    row = _note(
        ws,
        row,
        "Catatan: status perizinan harus diverifikasi pada OSS-RBA dan instansi sektor. Keberadaan NIB tidak otomatis berarti seluruh izin turunan aktif atau kewajiban pemenuhan telah dipenuhi.",
        13,
    )
    _finish_sheet(ws, row, 13)

    # Sheet 4 -----------------------------------------------------------------
    ws = wb.create_sheet(SHEET_COVENANTS)
    _title(ws, "NEGATIVE COVENANTS & UTANG BANK", "Fasilitas kredit, saldo, jaminan, dan pembatasan kreditur", 12, meta)
    row = 5
    row = _section(ws, row, "A. RINGKASAN UTANG BANK", 12)
    total_principal = sum(item.principal or 0 for item in analysis.loans)
    total_balance = sum(item.balance or 0 for item in analysis.loans)
    row = _kv_rows(
        ws,
        row,
        [
            ("Jumlah Fasilitas / Perjanjian", len(analysis.loans)),
            ("Total Plafon (Rp)", f"{total_principal:,.0f}" if total_principal else ""),
            ("Total Saldo Terutang (Rp)", f"{total_balance:,.0f}" if total_balance else ""),
            ("Negative Covenant Teridentifikasi", len(analysis.covenants)),
            ("Covenant HIGH", sum(1 for item in analysis.covenants if item.severity == "HIGH")),
        ],
    )
    row += 1
    row = _section(ws, row, "B. FASILITAS KREDIT & UTANG BANK", 12)
    loan_rows = [
        [
            index,
            item.document,
            item.creditor,
            item.facility,
            item.principal,
            item.balance,
            item.currency,
            item.due,
            item.restriction,
            item.status,
            item.severity,
            item.source,
        ]
        for index, item in enumerate(analysis.loans, start=1)
    ] or [[1, "Tidak ditemukan perjanjian kredit", "", "", "", "", "", "", "", "Perlu Klarifikasi", "HIGH", "Minta seluruh perjanjian kredit dan addendum."]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Perjanjian / Fasilitas", "Kreditur", "Jenis Fasilitas", "Plafon (Rp)", "Saldo (Rp)", "Mata Uang", "Jatuh Tempo", "Restriksi Utama", "Kepatuhan", "Severity", "Sumber"],
        loan_rows,
        widths=[5, 30, 20, 18, 16, 16, 9, 14, 35, 16, 10, 26],
        money_cols=(5, 6),
        severity_col=11,
        status_col=10,
    )
    row += 1
    row = _section(ws, row, "C. MATRIKS NEGATIVE COVENANTS", 12)
    covenant_rows = [
        [
            index,
            item.category,
            item.document,
            item.creditor,
            item.clause,
            item.restriction or item.summary,
            item.threshold,
            item.actual,
            item.status,
            item.severity,
            item.action,
            item.source,
        ]
        for index, item in enumerate(analysis.covenants, start=1)
    ] or [[1, "", "", "", "", "Tidak ada negative covenant yang dilaporkan.", "", "", "Perlu Klarifikasi", "MEDIUM", "Baca seluruh perjanjian kredit dan addendum.", ""]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Kategori", "Perjanjian", "Kreditur", "Pasal / Klausul", "Restriksi / Ringkasan", "Ambang / Batas", "Kondisi Aktual", "Status", "Severity", "Tindakan", "Sumber"],
        covenant_rows,
        widths=[5, 20, 26, 18, 15, 45, 18, 25, 16, 10, 32, 26],
        severity_col=10,
        status_col=9,
    )
    row = _note(
        ws,
        row,
        "Catatan: uji kepatuhan covenant terhadap kondisi aktual (perubahan kepemilikan, dividen, jaminan aset, utang tambahan, cross default). Pelanggaran dapat memicu percepatan jatuh tempo; pertimbangkan waiver atau consent kreditur sebelum closing.",
        12,
    )
    _finish_sheet(ws, row, 12)

    # Sheet 5 -----------------------------------------------------------------
    ws = wb.create_sheet(SHEET_LITIGATION)
    _title(ws, "LITIGASI & SENGKETA ASET", "Perkara berjalan, potensi klaim, dan status kepemilikan aset", 11, meta)
    row = 5
    row = _section(ws, row, "A. RINGKASAN LITIGASI & ASET", 11)
    litigation_value = sum(item.value or 0 for item in analysis.litigation)
    encumbered = sum(1 for item in analysis.assets if _text(item.encumbrance) and _text(item.encumbrance).lower() not in {"tidak ada", "none", "tidak ditemukan"})
    row = _kv_rows(
        ws,
        row,
        [
            ("Jumlah Perkara / Sengketa", len(analysis.litigation)),
            ("Nilai Gugatan / Klaim (Rp)", f"{litigation_value:,.0f}" if litigation_value else ""),
            ("Perkara Berjalan", sum(1 for item in analysis.litigation if item.status == "Berjalan")),
            ("Aset Tercatat", len(analysis.assets)),
            ("Aset dengan Indikasi Beban / Jaminan", encumbered),
        ],
    )
    row += 1
    row = _section(ws, row, "B. PERKARA LITIGASI & SENGKETA", 11)
    litigation_rows = [
        [
            index,
            item.case_number,
            item.court,
            item.parties,
            item.role,
            item.subject,
            item.value,
            item.status,
            item.severity,
            item.asset_impact,
            item.source,
        ]
        for index, item in enumerate(analysis.litigation, start=1)
    ] or [[1, "", "", "", "", "Tidak ditemukan perkara litigasi dalam data room.", "", "Tidak Ditemukan", "LOW", "", "Konfirmasi ada/tidaknya perkara."]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Nomor Perkara", "Pengadilan / Forum", "Para Pihak", "Peran Target", "Pokok Sengketa", "Nilai (Rp)", "Status", "Severity", "Dampak Aset", "Sumber"],
        litigation_rows,
        widths=[5, 24, 26, 30, 16, 35, 16, 15, 10, 30, 26],
        money_cols=(7,),
        severity_col=9,
        status_col=8,
    )
    row += 1
    row = _section(ws, row, "C. ASET & BEBAN JAMINAN", 11)
    asset_rows = [
        [
            index,
            item.name,
            item.asset_type,
            item.ownership,
            item.certificate,
            item.encumbrance,
            item.dispute,
            item.status,
            item.severity,
            item.source,
        ]
        for index, item in enumerate(analysis.assets, start=1)
    ] or [[1, "Tidak ditemukan dokumen aset", "", "", "", "Perlu Klarifikasi", "Perlu Klarifikasi", "Perlu Klarifikasi", "MEDIUM", "Minta sertifikat dan cek beban hak tanggungan."]]
    row, _ = _rows(
        ws,
        row,
        ["No", "Nama / Objek Aset", "Jenis Aset", "Kepemilikan", "Sertifikat / Bukti", "Beban / Jaminan", "Sengketa", "Status", "Severity", "Sumber"],
        asset_rows,
        widths=[5, 30, 18, 24, 24, 28, 20, 16, 10, 26],
        severity_col=9,
        status_col=8,
    )
    row = _note(
        ws,
        row,
        "Catatan: verifikasi keaslian sertifikat, riwayat kepemilikan, beban hak tanggungan/fidusia, dan status sengketa pada instansi berwenang. Daftar ini tidak menggantikan pengecekan sertifikat dan uji tuntas lapangan.",
        11,
    )
    _finish_sheet(ws, row, 11)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    summary = {
        "documents": len(analysis.documents),
        "shareholders": len(analysis.shareholders),
        "management": len(analysis.directors) + len(analysis.commissioners),
        "permits": len(analysis.permits),
        "kbli": len(analysis.kbli),
        "loans": len(analysis.loans),
        "covenants": len(analysis.covenants),
        "litigation": len(analysis.litigation),
        "assets": len(analysis.assets),
        "redFlags": {
            "total": len(analysis.red_flags),
            **severity_counts,
        },
        "overallRisk": _overall_risk(analysis),
    }
    return {
        "ok": True,
        "output": str(output),
        "sheets": [SHEET_SUMMARY, SHEET_OWNERSHIP, SHEET_PERMITS, SHEET_COVENANTS, SHEET_LITIGATION],
        "summary": summary,
        "warnings": getattr(analysis, "_warnings", []),
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _load_input(input_value: str) -> tuple[Any, str, Path | None]:
    if input_value == "-":
        try:
            return json.load(sys.stdin), "json", None
        except json.JSONDecodeError as exc:
            raise DueDiligenceError(f"Input JSON dari stdin tidak valid: {exc}") from exc
    path = Path(input_value)
    if not path.exists():
        raise DueDiligenceError(f"Input tidak ditemukan: {path}")
    if path.is_dir():
        return None, "directory", path
    if path.suffix.lower() != ".json":
        raise DueDiligenceError("Input file harus berupa JSON analisis atau direktori data room.")
    try:
        return json.loads(path.read_text(encoding="utf-8")), "json", path
    except (OSError, json.JSONDecodeError) as exc:
        raise DueDiligenceError(f"Input JSON tidak dapat dibaca: {exc}") from exc


def _default_output(input_value: str, input_type: str) -> Path:
    if input_value == "-":
        return Path("due-diligence-matrix.xlsx")
    path = Path(input_value)
    if input_type == "directory":
        return path.with_name(f"{path.name}.due-diligence.xlsx")
    return path.with_suffix(".due-diligence.xlsx")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build a five-sheet legal due diligence workbook from a data room or analysis JSON."
    )
    parser.add_argument("--input", required=True, help="Data room directory, analysis JSON, or - for stdin.")
    parser.add_argument("--output", default=None, help="Output .xlsx path.")
    parser.add_argument("--client", default="", help="Client name.")
    parser.add_argument("--target", default="", help="Target company name.")
    parser.add_argument("--reviewer", default="", help="Reviewer name.")
    parser.add_argument("--json", action="store_true", help="Print a machine-readable report.")
    args = parser.parse_args(argv)

    try:
        raw, input_type, input_path = _load_input(args.input)
        if input_type == "directory":
            assert input_path is not None
            analysis = analyse_directory(
                input_path,
                client=args.client,
                target=args.target,
                reviewer=args.reviewer,
            )
        else:
            analysis = normalise_analysis(raw)
            if args.client:
                analysis.deal.client = args.client
            if args.target:
                analysis.deal.target = args.target
                if not analysis.company.name:
                    analysis.company.name = args.target
            if args.reviewer:
                analysis.deal.reviewer = args.reviewer
            if not analysis.deal.date:
                analysis.deal.date = date.today().isoformat()
        output = Path(args.output) if args.output else _default_output(args.input, input_type)
        if output.suffix.lower() != ".xlsx":
            raise DueDiligenceError("Output harus berupa file .xlsx.")
        report = build_workbook(analysis, output)
        report["input"] = args.input if input_type == "directory" else str(input_path or "stdin")
        report["inputType"] = input_type
    except (DueDiligenceError, ValidationError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(report, ensure_ascii=False))
    else:
        summary = report["summary"]
        print(f"Due diligence workbook written to {report['output']}")
        print(
            f"{summary['documents']} documents, {summary['permits']} permits, "
            f"{summary['covenants']} covenants, {summary['litigation']} disputes; "
            f"overall risk {summary['overallRisk']}."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
